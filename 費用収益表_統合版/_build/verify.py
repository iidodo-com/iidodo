"""配布ブックを、元の3ファイルとは独立した経路で検証する。

  A. 旧・費用収益表（xls の計算済み値）とのセル単位の一致
  B. 経費区分の按分（旧・管理簿の計算済み値との一致）
  C. 台帳から Python で独立に再計算した月別・科目別の金額との一致（9月締め／10月締め）
  D. 数式エラーの有無

使い方: python3 verify.py <配布ブック.xlsx> <管理簿.xlsx> <貸付予算差引簿.xls> <費用収益表.xls>
"""
import collections
import math
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl
import xlrd

warnings.filterwarnings('ignore')
sys.path.insert(0, str(Path(__file__).parent))
import build as B
import extract as X


def xround(v):
    return math.floor(abs(v) + 0.5) * (1 if v >= 0 else -1)


def split(a, k):
    if k == 1:
        i, j, kk = xround(a * 10 / 23), xround(a / 23), xround(a / 23)
        return a - i - j - kk, i, j, kk
    if k == 2:
        return a, 0, 0, 0
    if k == 3:
        j = kk = xround(a / 12)
        return 0, a - j - kk, j, kk
    if k == 4:
        return 0, a, 0, 0
    if k == 5:
        i = xround(a * 10 / 11)
        return 0, i, a - i, 0
    if k == 6:
        return 0, 0, 0, a
    if k == 7:
        i = xround(a * 10 / 11)
        return 0, i, 0, a - i
    if k == 8:
        return 0, 0, a, 0


def recalc_with_month(final, month_label):
    tmp = Path(tempfile.mkdtemp())
    wb = openpyxl.load_workbook(final)
    wb['月次締め']['C5'] = month_label
    f = tmp / 'v.xlsx'
    wb.save(f)
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir',
                    str(tmp / 'o'), str(f)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return openpyxl.load_workbook(tmp / 'o' / 'v.xlsx', data_only=True)


def main(final, gyomu, kashi, hiyou):
    res = {}
    new = openpyxl.load_workbook(final, data_only=True)
    old = xlrd.open_workbook(hiyou)
    # D. エラー
    errs = [(ws.title, c.coordinate, c.value) for ws in new for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value[:2] in ('#V', '#N', '#R', '#D')]
    res['D.数式エラー'] = len(errs)
    # A. 旧表とのセル単位の一致
    tot = bad = 0
    for oldname, newname in B.OLD_NAMES.items():
        s = old.sheet_by_name(oldname)
        ws = new[newname]
        merged = set()
        for mr in ws.merged_cells.ranges:
            for rr in range(mr.min_row, mr.max_row + 1):
                for cc in range(mr.min_col, mr.max_col + 1):
                    if (rr, cc) != (mr.min_row, mr.min_col):
                        merged.add((rr, cc))
        for r in range(4, 63):
            for c in range(3, 22):
                if (r, c) in merged:
                    continue
                ov = s.cell_value(r - 1, c - 1) if r - 1 < s.nrows and c - 1 < s.ncols else ''
                ov = ov if isinstance(ov, float) else 0.0
                nv = ws.cell(r, c).value
                nv = nv if isinstance(nv, (int, float)) else 0.0
                tot += 1
                if abs(ov - nv) > 0.5:
                    bad += 1
                    print('A不一致', newname, r, c, ov, nv)
    res['A.旧表とのセル比較'] = f'{tot} セル中 不一致 {bad}'
    # B. 按分
    og = openpyxl.load_workbook(gyomu, data_only=True)
    led = new[B.SH_IN_G]
    n = bad = 0
    for r in range(B.G_FIRST, B.G_LAST + 1):
        src = led.cell(r, 18).value
        if not src:
            continue
        sh, rr = src.split('!')
        o = og[sh]
        ov = [round(o.cell(int(rr), c).value or 0) for c in (8, 9, 10, 11)]
        nv = [round(led.cell(r, c).value or 0) for c in (10, 11, 12, 13)]
        n += 1
        bad += ov != nv
    res['B.経費区分の按分（旧管理簿の計算値と比較）'] = f'{n} 行中 不一致 {bad}'
    # C. 独立再計算
    g, _ = X.load_gyomu(gyomu)
    k, _ = X.load_kashitsuke(kashi)
    for r in k:
        r['key'] = B.kashi_key(r)
    _, oldvals = X.load_old_hiyou(hiyou, tempfile.mkdtemp())
    B.migrate_residuals(k, oldvals['R8貸付'])
    kd = collections.defaultdict(float)
    for r in k:
        if r['st'] == '確認済' and r['kind'] == '通常':
            kd[(r['key'], (r['date'].month - 4) % 12 + 1)] += r['amt']
    for label in ('9月', '10月'):
        cut = int(label[:-1]) - 3
        wb = recalc_with_month(final, label)
        tot = bad = 0
        for sh, side in ((B.SH_TAN, '短'), (B.SH_CHO, '長')):
            ex = collections.defaultdict(float)
            for r in g:
                if r['st'] != '確認済':
                    continue
                key = B.GYOMU_MAP[f"{r['chu']}|{r['sho'] or ''}"]
                if key == '対象外':
                    continue
                idx = 13 if r['kind'] == '決算整理' else (r['date'].month - 4) % 12 + 1
                t, h, ke, ta = split(r['amt'], r['kbn'])
                ex[(key, idx)] += t if side == '短' else h + ke + ta
            ws = wb[sh]
            for row, key in B.GYOMU_ROWS.items():
                if row in (56, 57):
                    continue
                for idx in range(1, cut + 1):
                    tot += 1
                    if abs((ws.cell(row, 2 + idx).value or 0) - ex.get((key, idx), 0)) > 0.5:
                        bad += 1
            for idx in range(1, cut + 1):
                tot += 1
                v = (ws.cell(56, 2 + idx).value or 0) + (ws.cell(57, 2 + idx).value or 0)
                if abs(v - ex.get(('非常勤職員手当', idx), 0)) > 0.5:
                    bad += 1
        ws = wb[B.SH_KASHI]
        for row, key in B.KASHI_ROWS.items():
            for idx in range(1, cut + 1):
                tot += 1
                if abs((ws.cell(row, 2 + idx).value or 0) - kd.get((key, idx), 0)) > 0.5:
                    bad += 1
        res[f'C.独立再計算（{label}締め）'] = f'{tot} セル中 不一致 {bad}／検算シートのNG件数 {wb["検算"]["C3"].value}'
    for k_, v in res.items():
        print(f'{k_}: {v}')
    return res


if __name__ == '__main__':
    main(*sys.argv[1:5])
