"""旧3ファイル（業務経理区分別執行管理簿 / 貸付予算差引簿 / 費用収益表）から台帳データを取り出す。

状態（旧: 文字色・塗りつぶし）の判定ルール
  赤字                         -> 昨年実績
  黒字（青塗りあり）            -> 確認済
  黒字（塗りなし）              -> 未確認
  ただし、旧・費用収益表が既に計上している月（BASE_MONTH 以前）の黒字は、
  青塗りの有無にかかわらず 確認済 として移行する（旧表の数値と一致させるため）。
"""
import datetime
import re
import subprocess
import tempfile
from pathlib import Path

import openpyxl
import xlrd

BASE_MONTH = 9          # 旧・費用収益表で黒字転記済みの最終月（4月=4 ... 9月=9）
FY_START = datetime.date(2026, 4, 1)
FY_END = datetime.date(2027, 3, 31)

RED = 'FFFF0000'


def _is_red(cell):
    fc = cell.font.color
    return fc is not None and fc.type == 'rgb' and fc.rgb == RED


def _is_blue_fill(cell):
    if not cell.fill.fill_type:
        return False
    fg = cell.fill.fgColor
    return fg is not None and fg.type == 'theme' and fg.theme == 8


def load_gyomu(path):
    """管理簿(xlsx) の 月例 / 決算整理 を台帳行にする。"""
    wb = openpyxl.load_workbook(path)
    rows = []
    notes = {'promoted': [], 'kessan_black': []}
    for sheet, last in (('月例', 788), ('決算整理', 1835)):
        ws = wb[sheet]
        for r in range(5, last + 1):
            amt = ws.cell(r, 6).value
            d = ws.cell(r, 1).value
            if amt is None and d is None:
                continue
            red = _is_red(ws.cell(r, 6))
            blue = _is_blue_fill(ws.cell(r, 6))
            date = d.date() if isinstance(d, datetime.datetime) else d
            month = date.month if isinstance(date, datetime.date) else None
            chu = ws.cell(r, 3).value
            sho = ws.cell(r, 4).value
            tek = ws.cell(r, 5).value
            kbn = ws.cell(r, 7).value
            if sheet == '決算整理':
                # 令和7年度の決算整理（令和8年4月支払）は昨年実績として移行
                st = '昨年実績'
                if not red:
                    notes['kessan_black'].append((sheet, r, date, tek, amt))
            elif red:
                st = '昨年実績'
            elif month is not None and 4 <= month <= BASE_MONTH:
                st = '確認済'
                if not blue:
                    notes['promoted'].append((sheet, r, date, tek, amt))
            else:
                st = '確認済' if blue else '未確認'
            rows.append(dict(src=f'{sheet}!{r}', kind=sheet, date=date, chu=chu, sho=sho,
                             tek=tek, amt=amt, kbn=kbn, st=st))
    return rows, notes


def load_kashitsuke(path):
    """貸付予算差引簿(xls) の各ブロックを台帳行にする。"""
    b = xlrd.open_workbook(path, formatting_info=True)

    def info(s, r, c):
        xf = b.xf_list[s.cell_xf_index(r, c)]
        fc = b.font_list[xf.font_index].colour_index
        bg = xf.background.pattern_colour_index if xf.background.fill_pattern else None
        return fc, bg

    rows = []
    notes = {'fixed_date': [], 'promoted': [], 'nodate': []}
    for si, s in enumerate(b.sheets()):
        for c0, hc in ((0, 1), (5, 6)):
            hdrs = [r for r in range(s.nrows) if s.cell_value(r, c0 + 3) == '貸付経理']
            for i, h in enumerate(hdrs):
                end = hdrs[i + 1] if i + 1 < len(hdrs) else s.nrows
                block = s.cell_value(h, hc)
                notes.setdefault('blocks', []).append(dict(sheet=si, c0=c0, label=block, header=h + 1,
                                                           first=h + 3, last=end))
                for r in range(h + 2, end):
                    d, t, a = s.cell_value(r, c0), s.cell_value(r, c0 + 1), s.cell_value(r, c0 + 2)
                    if t == '' or not isinstance(a, float):
                        continue
                    if not (isinstance(d, float) or (isinstance(d, str) and re.fullmatch(r'\d+/', d))):
                        continue            # 「支出済額」などの集計行
                    fc, bg = info(s, r, c0 + 2)
                    _, bgd = info(s, r, c0)
                    red = fc == 10
                    blue = bg == 27
                    yellow = bgd == 13 or info(s, r, c0 + 1)[1] == 13
                    date = xlrd.xldate_as_datetime(d, 0).date() if isinstance(d, float) else None
                    month = date.month if date else None
                    if date is None:
                        notes['nodate'].append((si, r + 1, t, a))
                    if red:
                        st = '昨年実績'
                    elif month is not None and 4 <= month <= BASE_MONTH:
                        st = '確認済'
                        if not blue:
                            notes['promoted'].append((si, r + 1, date, t, a))
                    else:
                        st = '確認済' if blue else '未確認'
                    if (not red) and date and not (FY_START <= date <= FY_END):
                        fixed = date.replace(year=date.year + 1)
                        notes['fixed_date'].append((si, r + 1, date, fixed, t, a))
                        date = fixed
                    rows.append(dict(src=f"管理簿{'①' if si == 0 else '②'}!{r + 1}", sheet=si, block=block, rowno=r + 1, c0=c0,
                                     date=date, tek=t, amt=a, st=st,
                                     kind='決算整理' if yellow else '通常'))
    return rows, notes


def convert_xls_to_xlsx(xls_path, outdir):
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', str(outdir), str(xls_path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return Path(outdir) / (Path(xls_path).stem + '.xlsx')


def load_old_hiyou(xls_path, workdir):
    """旧・費用収益表: 数式(変換xlsx) と 計算済み値(xls) の両方を返す。"""
    conv = convert_xls_to_xlsx(xls_path, workdir)
    b = xlrd.open_workbook(xls_path)
    values = {}
    for s in b.sheets():
        vals = {}
        for r in range(s.nrows):
            for c in range(s.ncols):
                v = s.cell_value(r, c)
                if v != '':
                    vals[(r + 1, c + 1)] = v
        values[s.name] = vals
    return conv, values
