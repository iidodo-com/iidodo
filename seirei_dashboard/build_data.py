#!/usr/bin/env python3
"""data/raw/ の取得物を整形し、data/dataset.json を作る（HTML生成・テストの入力）。

取得できなかった／読み取れなかった値は null + status で記録し、補完・推測はしない。
  status: ok      … 資料に数値あり
          dash    … 資料上のハイフン表記（値なし）。順位対象外
          missing … 未取得（取得失敗・環境変数なし・読み取り不可）
"""
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

import sources as S

warnings.simplefilter("ignore")  # openpyxl の DrawingML 警告

ROOT = Path(__file__).resolve().parent
RAW = ROOT / "data" / "raw"
MANIFEST = ROOT / "data" / "manifest.json"
OUT = ROOT / "data" / "dataset.json"

CITY_CODES = [c[0] for c in S.CITIES]
CITY_NAME = {c[0]: c[1] for c in S.CITIES}
CITY_PREF = {c[0]: c[2] for c in S.CITIES}


def norm(s):
    return re.sub(r"\s+", "", str(s)) if s is not None else ""


def cell(value, status="ok", **extra):
    return {"value": value, "status": status, **extra}


def missing(reason):
    return {"value": None, "status": "missing", "reason": reason}


def blank_values(reason):
    return {c: missing(reason) for c in CITY_CODES}


def num(v):
    """資料の値を数値化。数値ならそのまま、'-' 等は None。"""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    return None


# ---------------------------------------------------------------- 定員管理調査
STAFF_COLS = ["一般管理", "福祉関係", "一般行政計", "教育", "警察", "消防", "普通会計計", "公営企業等会計", "合計"]


def parse_teiin(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = {}
    for r in ws.iter_rows(values_only=True):
        v = [c for c in r if c is not None]
        if v and isinstance(v[0], str) and len(v) == 10:
            rows[norm(v[0])] = v[1:]
    cities, total = {}, None
    for name, vals in rows.items():
        if name == "合計":
            total = dict(zip(STAFF_COLS, vals))
        elif name in {CITY_NAME[c] for c in CITY_CODES}:
            cities[name] = dict(zip(STAFF_COLS, vals))
    return cities, total


# ---------------------------------------------------------------- 住民基本台帳人口
def parse_jyuki(path):
    """市区町村別【総計】。政令市の行と、直後に続く区の行を読む。"""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = [r for r in ws.iter_rows(min_row=7, values_only=True) if r[2]]
    out = {}
    for code in CITY_CODES:
        name, pref = CITY_NAME[code], CITY_PREF[code]
        idx = [i for i, r in enumerate(rows) if norm(r[2]) == name and norm(r[1]) == pref]
        if len(idx) != 1:
            continue
        i = idx[0]
        city = rows[i]
        wards, skipped = [], []
        for r in rows[i + 1:]:
            nm = norm(r[2])
            if nm.startswith(name) and nm != name:
                # 区再編前の旧区行（値が *** で秘匿されている）は区合計の突合から除く
                (skipped if "再編前" in nm else wards).append(r)
            else:
                break
        out[code] = {
            "male": city[3], "female": city[4], "total": city[5], "households": city[6],
            "ward_sum": {"male": sum(w[3] for w in wards), "female": sum(w[4] for w in wards),
                         "total": sum(w[5] for w in wards), "households": sum(w[6] for w in wards)},
            "ward_count": len(wards), "skipped_pre_reorg_rows": [norm(s[2]) for s in skipped],
            "dantai_code": str(city[0]),
        }
    return out


# ---------------------------------------------------------------- 国勢調査(e-Stat API)
def parse_census(path):
    d = json.loads(Path(path).read_text(encoding="utf-8"))["GET_STATS_DATA"]
    vals = d["STATISTICAL_DATA"]["DATA_INF"]["VALUE"]
    out = {}
    for v in vals:
        area, sex = v["@area"], v["@cat01"]
        out.setdefault(area, {})[{"0": "total", "1": "male", "2": "female"}[sex]] = int(v["$"])
    return out


# ---------------------------------------------------------------- 財政状況資料集 総括表
FIN_LABELS = {
    "財政力指数": ("fin_zaiseiryoku", "財政力指数", "", 2),
    "経常収支比率": ("fin_keijo", "経常収支比率", "％", 1),
    "実質収支比率": ("fin_jissitsu_shushi", "実質収支比率", "％", 1),
    "実質公債費比率": ("fin_jissitsu_kousai", "実質公債費比率", "％", 1),
    "将来負担比率": ("fin_shorai", "将来負担比率", "％", 1),
    "標準財政規模": ("fin_hyojun", "標準財政規模", "千円", 0),
    "歳入総額": ("fin_saisyutsu_in", "歳入総額", "千円", 0),
    "歳出総額": ("fin_saisyutsu", "歳出総額", "千円", 0),
    "地方債現在高": ("fin_chihosai", "地方債現在高", "千円", 0),
    "基準財政収入額": ("fin_kijun_shunyu", "基準財政収入額", "千円", 0),
    "基準財政需要額": ("fin_kijun_juyo", "基準財政需要額", "千円", 0),
}


def parse_zaisei(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["総括表"]
    found, extra = {}, {}
    for row in ws.iter_rows():
        for ci, c in enumerate(row):
            if c.value is None:
                continue
            t = norm(c.value).replace("(※1)", "").replace("（※1）", "")
            right = [x.value for x in row[ci + 1:] if x.value not in (None, "")]
            if t in FIN_LABELS and t not in found and right:
                found[t] = right[0]  # 左から1番目が令和5年度、2番目が令和4年度
            if t == "令06.01.01(人)" and "pop_r6" not in extra and right:
                extra["pop_r6"] = right[0]
            if t == "令和2年国調(人)" and "pop_census" not in extra and right:
                extra["pop_census"] = right[0]
            if t == "市町村名" and "city" not in extra and right:
                extra["city"] = norm(right[0])
    wb.close()
    return found, extra


# ---------------------------------------------------------------- 組み立て
def main():
    if not MANIFEST.exists():
        sys.exit("data/manifest.json がありません。先に fetch_data.py を実行してください。")
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    msrc = {s["id"]: s for s in manifest["sources"]}
    indicators, integrity = [], {}

    def src_ok(sid):
        return msrc.get(sid, {}).get("status") == "ok"

    def reason(sid):
        return msrc.get(sid, {}).get("reason") or "取得物なし"

    # 財政状況資料集（市別xlsx）は一度だけ読む
    fin, fin_extra = {}, {}
    if src_ok("zaisei"):
        for code in CITY_CODES:
            fin[code], fin_extra[code] = parse_zaisei(RAW / "zaisei" / f"zaisei_{code}.xlsx")

    # ---- 住基人口 R6 / R7
    jy = {}
    for sid in ("jyuki_r6", "jyuki_r7"):
        jy[sid] = parse_jyuki(RAW / sid / f"{sid}.xlsx") if src_ok(sid) else {}
    integrity["jyuki"] = {sid: {c: v for c, v in d.items()} for sid, d in jy.items()}

    def jy_values(sid, key="total"):
        d = jy[sid]
        return {c: (cell(d[c][key]) if c in d else missing(reason(sid) if not src_ok(sid) else "資料に該当行なし"))
                for c in CITY_CODES}

    pop_r6 = jy_values("jyuki_r6")
    pop_r7 = jy_values("jyuki_r7")
    indicators.append({
        "id": "pop_r6", "label": "人口（住民基本台帳・令和6年1月1日）", "group": "人口", "unit": "人", "kind": "raw",
        "sources": ["jyuki_r6"], "basis": "令和6年1月1日", "decimals": 0,
        "note": "住民基本台帳に基づく人口（日本人住民＋外国人住民）。人口あたり指標の分母に使用。",
        "values": pop_r6})
    indicators.append({
        "id": "pop_r7", "label": "人口（住民基本台帳・令和7年1月1日）", "group": "人口", "unit": "人", "kind": "raw",
        "sources": ["jyuki_r7"], "basis": "令和7年1月1日", "decimals": 0,
        "note": "参考。人口あたり指標には使用しない（職員数・決算と年が離れるため）。",
        "values": pop_r7})

    # 人口増減率 R6→R7（派生）
    chg = {}
    for c in CITY_CODES:
        a, b = pop_r6[c], pop_r7[c]
        if a["status"] == "ok" and b["status"] == "ok":
            chg[c] = cell(round((b["value"] - a["value"]) / a["value"] * 100, 4),
                          numerator=b["value"] - a["value"], denominator=a["value"],
                          inputs={"pop_r7": b["value"], "pop_r6": a["value"]})
        else:
            chg[c] = missing("元データ未取得")
    indicators.append({
        "id": "pop_change", "label": "人口増減率（令和6.1.1→令和7.1.1）", "group": "人口", "unit": "％", "kind": "derived",
        "sources": ["jyuki_r6", "jyuki_r7"], "basis": "令和6年1月1日→令和7年1月1日", "decimals": 2,
        "formula": "（令和7.1.1人口 − 令和6.1.1人口）÷ 令和6.1.1人口 × 100",
        "numerator_label": "人口増減数（人）", "denominator_label": "令和6.1.1人口（人）", "multiplier": 100,
        "inputs": ["pop_r7", "pop_r6"], "note": "2時点の差を取った指標。時点が異なることを承知の上で表示。",
        "values": chg})

    # ---- 国勢調査
    if src_ok("census2020"):
        cs = parse_census(RAW / "census2020" / "census2020.json")
        census_vals = {c: (cell(cs[c]["total"]) if c in cs and "total" in cs[c] else missing("API応答に該当なし"))
                       for c in CITY_CODES}
        integrity["census"] = cs
    else:
        census_vals = blank_values(reason("census2020"))
        integrity["census"] = {}
    # 財政状況資料集に記載の「令和2年国調」欄との差異を自動検出し、注記に残す（隠さない）
    census_note = "e-Stat API（統計表ID 0003445078）より取得した値。"
    diffs = {}
    if src_ok("census2020") and src_ok("zaisei"):
        for c in CITY_CODES:
            z = fin_extra.get(c, {}).get("pop_census")
            e = census_vals[c]["value"]
            if z is not None and e is not None and z != e:
                diffs[c] = {"estat": e, "zaisei": z}
        if diffs:
            census_note += " 財政状況資料集の国調欄との差異: " + "、".join(
                f"{CITY_NAME[c]}（e-Stat {d['estat']:,}人／資料集 {d['zaisei']:,}人）" for c, d in diffs.items()
            ) + "。本指標はe-Statの値を表示。"
    integrity["census_vs_zaisei_diff"] = diffs
    indicators.append({
        "id": "pop_census", "label": "人口（国勢調査・令和2年10月1日）", "group": "人口", "unit": "人", "kind": "raw",
        "sources": ["census2020"], "basis": "令和2年10月1日", "decimals": 0,
        "note": census_note,
        "values": census_vals})

    # ---- 職員数（定員管理調査）
    staff = {}
    teiin_total = None
    if src_ok("teiin"):
        staff, teiin_total = parse_teiin(RAW / "teiin" / "teiin_shitei.xlsx")
    integrity["teiin"] = {"cities": staff, "total_row": teiin_total}
    staff_defs = [
        ("staff_total", "職員数 合計（普通会計＋公営企業等会計）", "合計"),
        ("staff_general", "職員数 普通会計計", "普通会計計"),
        ("staff_admin", "職員数 一般行政部門計", "一般行政計"),
        ("staff_admin_mgmt", "職員数 一般行政：一般管理", "一般管理"),
        ("staff_welfare", "職員数 一般行政：福祉関係", "福祉関係"),
        ("staff_edu", "職員数 教育部門", "教育"),
        ("staff_fire", "職員数 消防部門", "消防"),
        ("staff_enterprise", "職員数 公営企業等会計", "公営企業等会計"),
    ]
    for iid, label, col in staff_defs:
        vals = {}
        for c in CITY_CODES:
            n = CITY_NAME[c]
            vals[c] = cell(staff[n][col]) if n in staff else missing(reason("teiin") if not src_ok("teiin") else "資料に該当行なし")
        indicators.append({
            "id": iid, "label": label, "group": "職員数", "unit": "人", "kind": "raw", "sources": ["teiin"],
            "basis": "令和6年4月1日", "decimals": 0,
            "note": "地方公共団体定員管理調査 第1表。警察部門は指定都市では0。", "values": vals})

    # ---- 人口1万人あたり職員数（派生）
    per10k = [
        ("staff_total", "人口1万人あたり職員数 合計"),
        ("staff_general", "人口1万人あたり職員数 普通会計計"),
        ("staff_admin", "人口1万人あたり職員数 一般行政部門"),
        ("staff_welfare", "人口1万人あたり職員数 福祉関係"),
        ("staff_edu", "人口1万人あたり職員数 教育部門"),
        ("staff_fire", "人口1万人あたり職員数 消防部門"),
    ]
    by_id = {i["id"]: i for i in indicators}
    for src_id, label in per10k:
        vals = {}
        for c in CITY_CODES:
            a, b = by_id[src_id]["values"][c], pop_r6[c]
            if a["status"] == "ok" and b["status"] == "ok":
                vals[c] = cell(round(a["value"] / b["value"] * 10000, 6), numerator=a["value"], denominator=b["value"],
                               inputs={src_id: a["value"], "pop_r6": b["value"]})
            else:
                vals[c] = missing("元データ未取得")
        indicators.append({
            "id": src_id + "_per10k", "label": label, "group": "職員数（人口あたり）", "unit": "人／1万人",
            "kind": "derived", "sources": ["teiin", "jyuki_r6"], "decimals": 2,
            "basis": "分子：令和6年4月1日／分母：令和6年1月1日（3か月の差）",
            "formula": f"{by_id[src_id]['label'].replace('職員数 ', '職員数（')}） ÷ 住基人口（令和6.1.1） × 10,000",
            "numerator_label": "職員数（人）", "denominator_label": "住基人口 令和6.1.1（人）", "multiplier": 10000,
            "inputs": [src_id, "pop_r6"],
            "note": "分子・分母の基準日が3か月ずれる。職員数は4月1日、人口は直前の1月1日の住基人口を採用。",
            "values": vals})
        by_id[indicators[-1]["id"]] = indicators[-1]

    # ---- 財政（財政状況資料集 総括表）
    integrity["zaisei"] = {"fin": fin, "extra": fin_extra}
    fin_meta = {
        "財政力指数": ("財政力指数", "", 2, "基準財政収入額÷基準財政需要額の過去3か年平均（資料の表記に従う）。"),
        "経常収支比率": ("経常収支比率", "％", 1, "経常的な経費に経常一般財源がどれだけ充当されたかの比率。"),
        "実質収支比率": ("実質収支比率", "％", 1, "標準財政規模に対する実質収支の比率。"),
        "実質公債費比率": ("実質公債費比率", "％", 1, "資料の表記に従う（3か年平均）。"),
        "将来負担比率": ("将来負担比率", "％", 1, "資料上『－』の市は値なし（順位対象外）。"),
        "標準財政規模": ("標準財政規模", "千円", 0, ""),
        "歳入総額": ("歳入総額（一般会計等）", "千円", 0, "財政状況資料集 総括表の「歳入総額」。"),
        "歳出総額": ("歳出総額（一般会計等）", "千円", 0, "財政状況資料集 総括表の「歳出総額」。"),
        "地方債現在高": ("地方債現在高", "千円", 0, "臨時財政対策債を含む。"),
    }
    for key, (label, unit, dec, note) in fin_meta.items():
        iid = FIN_LABELS[key][0]
        vals = {}
        for c in CITY_CODES:
            if c not in fin or key not in fin[c]:
                vals[c] = missing(reason("zaisei") if not src_ok("zaisei") else "資料に該当項目なし")
                continue
            raw = fin[c][key]
            n = num(raw)
            vals[c] = cell(n) if n is not None else {"value": None, "status": "dash", "raw_text": str(raw)}
        indicators.append({
            "id": iid, "label": label, "group": "財政", "unit": unit, "kind": "raw", "sources": ["zaisei"],
            "basis": "令和5年度決算", "decimals": dec, "note": note, "values": vals})
        by_id[iid] = indicators[-1]

    # 住民一人あたり（派生）。分母は住基人口 令和6.1.1（令和5年度内の日付）。
    for src_id, label in (("fin_saisyutsu", "住民一人あたり歳出額"), ("fin_chihosai", "住民一人あたり地方債現在高")):
        vals = {}
        for c in CITY_CODES:
            a, b = by_id[src_id]["values"][c], pop_r6[c]
            if a["status"] == "ok" and b["status"] == "ok":
                vals[c] = cell(round(a["value"] * 1000 / b["value"], 4), numerator=a["value"], denominator=b["value"],
                               inputs={src_id: a["value"], "pop_r6": b["value"]})
            else:
                vals[c] = missing("元データ未取得")
        indicators.append({
            "id": src_id + "_percap", "label": label, "group": "財政（住民あたり）", "unit": "円／人", "kind": "derived",
            "sources": ["zaisei", "jyuki_r6"], "decimals": 0,
            "basis": "分子：令和5年度決算／分母：令和6年1月1日人口（同年度内の日付）",
            "formula": f"{by_id[src_id]['label'].split('（')[0]}（千円） × 1,000 ÷ 住基人口（令和6.1.1）",
            "numerator_label": by_id[src_id]["label"].split("（")[0] + "（千円）", "denominator_label": "住基人口 令和6.1.1（人）",
            "multiplier": 1000, "inputs": [src_id, "pop_r6"],
            "note": "年度決算額を、年度内の1月1日人口で割る。", "values": vals})

    # ---- 未取得として明示する候補（今回の取得対象に含めていない）
    not_fetched = [
        ("地方公務員給与実態調査（平均給料月額・ラスパイレス指数 等）", "今回の取得対象に含めていない"),
        ("定員管理調査 第2表（細部門別職員数：税務・衛生・土木・病院・水道 等）", "今回の取得対象に含めていない（第1表の大分類のみ収録）"),
        ("決算状況調（目的別・性質別歳出の詳細）", "今回の取得対象に含めていない"),
    ]

    sources_out = []
    for sid in ("teiin", "jyuki_r6", "jyuki_r7", "census2020", "zaisei"):
        m = msrc.get(sid)
        if not m:
            continue
        lm = sorted({f.get("last_modified") for f in m["files"] if f.get("last_modified")})
        sources_out.append({
            "id": sid, "name": m["name"], "publisher": m["publisher"], "stat_id": m["stat_id"],
            "page_url": m["page_url"], "estat_url": m["estat_url"], "survey_date": m["survey_date"],
            "survey_label": m.get("survey_label"), "status": m["status"], "reason": m.get("reason"),
            "updated": m.get("updated") or (f"サーバー更新日: {lm[0]}" + (f" 〜 {lm[-1]}" if len(lm) > 1 else "") if lm else None),
            "retrieved_at": m["files"][0]["retrieved_at"] if m["files"] else None,
            "files": [{"name": f["name"], "url": f["url"], "sha256": f.get("sha256"), "status": f["status"]} for f in m["files"]],
        })

    dataset = {
        "title": "政令指定都市 比較ダッシュボード",
        "generated_from_fetch_at": manifest["fetched_at"],
        "highlight": S.HIGHLIGHT,
        "cities": [{"code": c[0], "name": c[1], "pref": c[2]} for c in S.CITIES],
        "sources": sources_out,
        "indicators": indicators,
        "not_fetched": [{"label": a, "reason": b} for a, b in not_fetched],
        "integrity": integrity,
    }
    OUT.write_text(json.dumps(dataset, ensure_ascii=False, indent=1), encoding="utf-8")
    nmiss = sum(1 for i in indicators for v in i["values"].values() if v["status"] == "missing")
    print(f"wrote {OUT}  指標{len(indicators)}件  未取得セル{nmiss}件")


if __name__ == "__main__":
    main()
