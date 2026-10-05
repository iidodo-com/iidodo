"""差分更新・中断再開・削除反映・安全装置・読み取り専用・FTS整合性のテスト。"""
import builtins
import io
import os
import shutil
import tempfile
import unittest
from unittest import mock

from tests import common
import generate_sample
import index as index_cli
import search as search_cli
from core import roots as R
import verify
from core import db, indexer, opener
from core.config import Config
from core.indexer import run_index
from core.kinds import KIND_CELLCAP, KIND_SIZE
from core.scanner import RootError
from core.searcher import build_options, search

DEL_REL = "契約書/秘密保持契約書.docx"
UPD_REL = "企画　資料/新規事業企画書.pptx"


def dump(conn):
    """索引の内容（ファイル・場所・本文）を、比較用の集合にして返す。"""
    return {(r[0], r[1], r[2]) for r in conn.execute(
        "SELECT f.relpath, c.location, c.text FROM chunks c JOIN files f ON f.id=c.file_id")}


def find(conn, term, scope="place"):
    """語で検索した (相対パス, 場所) の集合を返す。"""
    return {(r.relpath, r.location) for r in search(conn, term, build_options(limit=1000, scope=scope)).results}


class IncrementalTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.work, self.exp, self.cfg, self.conn = common.fresh_sample(self.tmp)
        self.addCleanup(self.conn.close)
        self.docs = self.cfg.roots[0]
        self.first = run_index(self.cfg, self.conn)

    def test_update_and_delete(self):
        """1つ更新・1つ削除して再実行 → 反映され、更新していないファイルは再処理されず、FTS索引の整合性も保たれる。"""
        self.assertIn(UPD_REL, self.first.processed)
        self.assertEqual(len(self.first.processed), self.first.total_targets)
        self.assertEqual(find(self.conn, "損害賠償"), {(DEL_REL, "段落6"), (UPD_REL, "スライド4"), ("報告/調査報告書.pdf", "p.3")})
        # 更新: 企画書に新しい語句を含むスライドを足して保存し直す（更新日時も変わる）
        generate_sample_pptx = os.path.join(self.docs, *UPD_REL.split("/"))
        from pptx import Presentation
        prs = Presentation(generate_sample_pptx)
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        from pptx.util import Inches
        sl.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "追加された新スライドの語句ゼータ"
        prs.save(generate_sample_pptx)
        os.utime(generate_sample_pptx, (1_800_000_000, 1_800_000_000))
        os.remove(os.path.join(self.docs, *DEL_REL.split("/")))
        calls = []
        real = indexer.extract
        with mock.patch.object(indexer, "extract", side_effect=lambda ext, data, cfg: (calls.append(1), real(ext, data, cfg))[1]):
            st = run_index(self.cfg, self.conn)
        self.assertEqual(st.processed, [UPD_REL])
        self.assertEqual(len(calls), 1)  # 抽出関数が呼ばれたのは更新した1ファイルだけ
        self.assertEqual(st.deleted, [DEL_REL])
        self.assertEqual(st.skipped_unchanged, st.total_targets - 1)
        self.assertEqual(find(self.conn, "ゼータ"), {(UPD_REL, "スライド5")})
        self.assertEqual(find(self.conn, "損害賠償"), {(UPD_REL, "スライド4"), ("報告/調査報告書.pdf", "p.3")})
        self.assertEqual(find(self.conn, "秘密保持契約書"), set())
        self.assertEqual(self.conn.execute("SELECT count(*) FROM files WHERE relpath=?", (DEL_REL,)).fetchone()[0], 0)
        ok, msg = db.check_integrity(self.conn)
        self.assertTrue(ok, msg)

    def test_scan_progress_is_reported(self):
        """走査中にも進捗（見つけたファイル数・調べているフォルダ）が通知され、数は増えていく。最初のファイル処理より前に出る。"""
        events, order = [], []
        run_index(self.cfg, self.conn, lambda *a: order.append("p"),
                  scan_progress=lambda found, d, el: (events.append((found, d)), order.append("s")))
        self.assertGreater(len(events), 5)
        counts = [e[0] for e in events]
        self.assertEqual(counts, sorted(counts))
        self.assertIn("契約書", [e[1] for e in events])
        self.assertLess(max(i for i, x in enumerate(order) if x == "s"), min(i for i, x in enumerate(order) if x == "p"))
        # 中止（例外）も、走査中に効く
        class Stop(Exception):
            pass

        def stop(found, d, el):
            raise Stop()
        with self.assertRaises(Stop):
            run_index(self.cfg, self.conn, scan_progress=stop)

    def test_no_change_rerun_reprocesses_nothing(self):
        """何も変えずに再実行すると、1ファイルも再処理されない（エラーのファイルも同様）。"""
        with mock.patch.object(indexer, "extract", side_effect=AssertionError("再処理された")):
            st = run_index(self.cfg, self.conn)
        self.assertEqual(st.processed, [])
        self.assertEqual(st.skipped_unchanged, st.total_targets)
        self.assertTrue(db.check_integrity(self.conn)[0])

    def test_retry_errors(self):
        """エラーのファイルは、変更がなければ再試行されず、--retry-errors で再試行される。"""
        n_err = self.conn.execute("SELECT count(*) FROM files WHERE status='error'").fetchone()[0]
        self.assertGreater(n_err, 0)
        self.assertEqual(run_index(self.cfg, self.conn).processed, [])
        st = run_index(self.cfg, self.conn, retry_errors=True)
        self.assertEqual(len(st.processed), n_err)
        self.assertTrue(db.check_integrity(self.conn)[0])

    def test_extract_setting_change_reprocesses_all(self):
        """抽出に影響する設定（xlsxのセル上限など）を変えたら、変更のないファイルも再処理される。"""
        self.cfg.max_cells_per_xlsx = 7
        st = run_index(self.cfg, self.conn)
        self.assertEqual(len(st.processed), st.total_targets)

    def test_resume_after_interrupt_between_files(self):
        """途中で中断（Ctrl+C）しても、次回は処理済みの分を飛ばして再開し、最終結果は通しで実行した場合と同じになる。"""
        full = dump(self.conn)
        conn2 = db.connect_rw(os.path.join(self.tmp, "resume.db"))
        self.addCleanup(conn2.close)

        def stop_at_5(done, total, elapsed, cur):
            if done == 5:
                raise KeyboardInterrupt()
        with self.assertRaises(KeyboardInterrupt):
            run_index(self.cfg, conn2, stop_at_5)
        n1 = conn2.execute("SELECT count(*) FROM files WHERE status!='excluded'").fetchone()[0]
        self.assertEqual(n1, 5)
        st = run_index(self.cfg, conn2)
        self.assertEqual(st.skipped_unchanged, 5)
        self.assertEqual(len(st.processed), st.total_targets - 5)
        self.assertEqual(dump(conn2), full)
        self.assertTrue(db.check_integrity(conn2)[0])

    def test_interrupt_inside_a_file_rolls_back(self):
        """1ファイルの処理中に中断しても、そのファイルは中途半端に登録されず、次回に処理される。"""
        conn2 = db.connect_rw(os.path.join(self.tmp, "resume2.db"))
        self.addCleanup(conn2.close)
        real = indexer.extract
        seen = []

        def boom(ext, data, cfg):
            seen.append(1)
            if len(seen) == 3:
                raise KeyboardInterrupt()
            return real(ext, data, cfg)
        with mock.patch.object(indexer, "extract", side_effect=boom):
            with self.assertRaises(KeyboardInterrupt):
                run_index(self.cfg, conn2)
        self.assertEqual(conn2.execute("SELECT count(*) FROM files WHERE status!='excluded'").fetchone()[0], 2)
        self.assertTrue(db.check_integrity(conn2)[0])
        st = run_index(self.cfg, conn2)
        self.assertEqual(len(st.processed), st.total_targets - 2)
        self.assertEqual(dump(conn2), dump(self.conn))

    def test_integrity_check_detects_corruption(self):
        """integrity-check は、索引と本文の不整合を実際に検出できる（テストが空振りでないことの確認）。"""
        self.assertTrue(db.check_integrity(self.conn)[0])
        self.conn.execute("INSERT INTO chunks(file_id, location, text) VALUES(1, 'x', '索引に載せていない本文')")
        self.conn.commit()
        ok, msg = db.check_integrity(self.conn)
        self.assertFalse(ok)
        self.assertIn("整合性エラー", msg)
        db.rebuild_fts(self.conn)
        self.assertTrue(db.check_integrity(self.conn)[0])

    def test_renamed_file_is_moved_in_index(self):
        """リネームされたファイルは、古い名前が削除され、新しい名前で登録される（元ファイルは検索側で触らない）。"""
        src = os.path.join(self.docs, "報告", "共通メモ.docx")
        dst = os.path.join(self.docs, "報告", "共通メモ（改名後）.docx")
        os.rename(src, dst)
        st = run_index(self.cfg, self.conn)
        self.assertEqual((st.processed, st.deleted), (["報告/共通メモ（改名後）.docx"], ["報告/共通メモ.docx"]))
        self.assertTrue(db.check_integrity(self.conn)[0])


class SafetyTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def make(self, n):
        """n個の小さなdocxを作って、取り込み済みの (cfg, conn) を返す。"""
        import docx
        root = os.path.join(self.tmp, "r")
        os.makedirs(root)
        for i in range(n):
            d = docx.Document()
            d.add_paragraph("ダミー文書%d番の本文" % i)
            d.save(os.path.join(root, "f%02d.docx" % i))
        cfg = Config(roots=[root], db_path=os.path.join(self.tmp, "i.db"), output_dir=os.path.join(self.tmp, "o"))
        conn = db.connect_rw(cfg.db_path)
        self.addCleanup(conn.close)
        run_index(cfg, conn)
        return cfg, conn

    def test_mass_delete_guard(self):
        """半数を超える削除は、安全装置で止まり、--allow-mass-delete で実行できる。"""
        cfg, conn = self.make(12)
        for i in range(8):
            os.remove(os.path.join(cfg.roots[0], "f%02d.docx" % i))
        st = run_index(cfg, conn)
        self.assertEqual((st.delete_blocked, st.deleted), (8, []))
        self.assertEqual(conn.execute("SELECT count(*) FROM files").fetchone()[0], 12)
        st = run_index(cfg, conn, allow_mass_delete=True)
        self.assertEqual(len(st.deleted), 8)
        self.assertEqual(conn.execute("SELECT count(*) FROM files").fetchone()[0], 4)
        self.assertTrue(db.check_integrity(conn)[0])

    def test_small_deletion_is_applied(self):
        """少数の削除は、安全装置にかからず反映される。"""
        cfg, conn = self.make(12)
        os.remove(os.path.join(cfg.roots[0], "f00.docx"))
        self.assertEqual(run_index(cfg, conn).deleted, ["f00.docx"])

    def test_unreachable_root_deletes_nothing(self):
        """対象フォルダに到達できないとき（共有フォルダ未接続など）は、何も削除せず中止する。"""
        cfg, conn = self.make(3)
        os.rename(cfg.roots[0], cfg.roots[0] + "_x")
        with self.assertRaises(RootError) as cm:
            run_index(cfg, conn)
        self.assertIn("アクセスできません", str(cm.exception))
        self.assertEqual(conn.execute("SELECT count(*) FROM files").fetchone()[0], 3)

    def test_unreadable_folder_keeps_entries(self):
        """読めないフォルダ配下のファイルは、「消えた」と判断せず索引に残す。走査エラーは記録される。"""
        cfg, conn = self.make(0)
        sub = os.path.join(cfg.roots[0], "sub")
        os.makedirs(sub)
        import docx
        d = docx.Document()
        d.add_paragraph("サブフォルダの本文")
        d.save(os.path.join(sub, "a.docx"))
        run_index(cfg, conn)
        self.assertEqual(find(conn, "サブフォルダの本文", "place"), {("sub/a.docx", "段落1")})
        real = os.scandir

        def deny(path):
            if os.path.basename(str(path)) == "sub":
                raise PermissionError(13, "Permission denied")
            return real(path)
        with mock.patch("core.scanner.os.scandir", side_effect=deny):
            st = run_index(cfg, conn)
        self.assertEqual(st.deleted, [])
        self.assertEqual(find(conn, "サブフォルダの本文", "place"), {("sub/a.docx", "段落1")})
        self.assertEqual(conn.execute("SELECT count(*) FROM issues WHERE kind='走査エラー'").fetchone()[0], 1)

    def test_size_limit(self):
        """サイズ上限を超えたファイルは「対象外(サイズ)」として記録され、本文は取り込まれない。"""
        cfg, conn = self.make(2)
        cfg.max_file_size_mb = 0.0005  # 約524バイト
        st = run_index(cfg, conn)
        n = conn.execute("SELECT count(*) FROM issues WHERE kind=?", (KIND_SIZE,)).fetchone()[0]
        self.assertEqual(n, 2)
        self.assertEqual(conn.execute("SELECT count(*) FROM chunks").fetchone()[0], 0)
        cfg.max_file_size_mb = 100  # 上限を戻せば、再び取り込まれる（対象外は毎回判定し直す）
        run_index(cfg, conn)
        self.assertEqual(conn.execute("SELECT count(*) FROM chunks").fetchone()[0], 2)
        self.assertTrue(db.check_integrity(conn)[0])

    def test_xlsx_cell_cap(self):
        """Excelのセル数上限: 超えた分は取り込まず、truncated と「セル上限」の記録が残る。"""
        import openpyxl
        root = os.path.join(self.tmp, "r")
        os.makedirs(root)
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in range(1, 3001):
            ws.cell(r, 1, "語句%d" % r)
        wb.save(os.path.join(root, "大きい表.xlsx"))
        cfg = Config(roots=[root], db_path=os.path.join(self.tmp, "i.db"), output_dir=self.tmp, max_cells_per_xlsx=1000)
        conn = db.connect_rw(cfg.db_path)
        self.addCleanup(conn.close)
        run_index(cfg, conn)
        self.assertEqual(conn.execute("SELECT count(*) FROM chunks").fetchone()[0], 1000)
        self.assertEqual(conn.execute("SELECT truncated FROM files").fetchone()[0], 1)
        self.assertEqual(conn.execute("SELECT count(*) FROM issues WHERE kind=?", (KIND_CELLCAP,)).fetchone()[0], 1)
        self.assertEqual(find(conn, "語句1000"), {("大きい表.xlsx", "Sheet!A1000")})
        self.assertEqual(find(conn, "語句1001"), set())


class ScopeTest(unittest.TestCase):
    """検索フォルダの選択: 選んだフォルダ以下だけを検索し、親は見ない。後から追加したフォルダ・入れ子のまとめ。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.work, self.exp, self.cfg, self.conn = common.fresh_sample(self.tmp)
        self.addCleanup(self.conn.close)
        self.docs = self.cfg.roots[0]
        run_index(self.cfg, self.conn)

    def names(self, folder, q="ABC", scope="file"):
        """folder（絶対パス）以下を検索した結果の相対パスの一覧。"""
        o = build_options(folder=folder, limit=1000, scope=scope, roots=R.indexed_roots(self.conn))
        return sorted(r.relpath for r in search(self.conn, q, o).results)

    def test_child_folder_only_not_parent_or_siblings(self):
        """子フォルダを選ぶと、その下だけを検索し、親フォルダ直下や兄弟フォルダのファイルは出ない。"""
        self.assertEqual(self.names(os.path.join(self.docs, "契約書")), ["契約書/業務委託契約書\u3000雛形.docx", "契約書/秘密保持契約書.docx"])
        self.assertEqual(self.names(os.path.join(self.docs, "報告")), ["報告/調査報告書.pdf"])
        self.assertEqual(self.names(os.path.join(self.docs, "企画\u3000資料")), ["企画\u3000資料/予算表 2024.xlsx", "企画\u3000資料/新規事業企画書.pptx"])
        self.assertEqual(len(self.names(self.docs)), 5)  # 親（root）を選べば全体

    def test_prefix_sibling_not_matched(self):
        """「契約書」を選んでも、名前が「契約書」で始まる別のフォルダ（契約書2）は含まれない。"""
        os.makedirs(os.path.join(self.docs, "契約書2"))
        import docx
        d = docx.Document()
        d.add_paragraph("ABC株式会社の別フォルダ文書")
        d.save(os.path.join(self.docs, "契約書2", "別.docx"))
        run_index(self.cfg, self.conn)
        self.assertNotIn("契約書2/別.docx", self.names(os.path.join(self.docs, "契約書")))
        self.assertEqual(self.names(os.path.join(self.docs, "契約書2")), ["契約書2/別.docx"])

    def test_deep_child_and_place_scope(self):
        """深い階層のフォルダ・場所単位の検索でも、選んだフォルダ以下に限られる。"""
        deep = "/".join(("長いフォルダ名" * 6)[:40] + str(i) for i in range(7))
        self.assertEqual(self.names(os.path.join(self.docs, *deep.split("/")[:3]), "長いパスの文書"), [deep + "/長いパスの文書.docx"])
        self.assertEqual(self.names(os.path.join(self.docs, "報告"), "損害賠償", "place"), ["報告/調査報告書.pdf"])

    def test_not_indexed_folder_is_japanese_error(self):
        """登録されていないフォルダを指定すると、日本語のエラー（作り方つき）になる。"""
        from core.query import QueryError
        with self.assertRaises(QueryError) as cm:
            build_options(folder=os.path.join(self.tmp, "別の場所"), roots=R.indexed_roots(self.conn))
        self.assertIn("まだインデックスされていません", str(cm.exception))

    def test_add_child_then_parent_absorbs_without_duplicates(self):
        """子フォルダを追加してから親フォルダを追加すると、子の登録は親にまとめられ、同じ文書が二重に出ない。"""
        conn = db.connect_rw(os.path.join(self.tmp, "add.db"))
        self.addCleanup(conn.close)
        child = os.path.join(self.docs, "契約書")
        self.assertEqual(R.add_root(conn, child, []), [])
        self.assertEqual(R.indexed_roots(conn), [])  # 作成が終わるまでは「未完了」
        run_index(self.cfg, conn, roots=[child])
        R.clear_incomplete(conn, [child])
        self.assertEqual(R.indexed_roots(conn), [child])
        n_child = conn.execute("SELECT count(*) FROM files WHERE root=?", (child,)).fetchone()[0]
        self.assertGreater(n_child, 0)
        absorbed = R.add_root(conn, self.docs, [])
        self.assertEqual(absorbed, [child])
        self.assertEqual(conn.execute("SELECT count(*) FROM files WHERE root=?", (child,)).fetchone()[0], 0)
        run_index(self.cfg, conn, roots=[self.docs])
        R.clear_incomplete(conn, [self.docs])
        self.assertEqual(R.indexed_roots(conn), [self.docs])
        self.assertEqual(R.load_extra_roots(conn), [self.docs])
        dup = conn.execute("SELECT count(*) FROM (SELECT relpath FROM files GROUP BY relpath HAVING count(*)>1)").fetchone()[0]
        self.assertEqual(dup, 0)
        self.assertEqual(dump(conn), dump(self.conn))
        self.assertTrue(db.check_integrity(conn)[0])
        # 子フォルダを選べば、親にまとめた後でも、その下だけが検索される
        o = build_options(folder=child, limit=1000, roots=R.indexed_roots(conn))
        self.assertEqual(sorted(r.relpath for r in search(conn, "ABC", o).results), ["契約書/業務委託契約書\u3000雛形.docx", "契約書/秘密保持契約書.docx"])

    def test_single_root_run_leaves_other_roots_alone(self):
        """1つのフォルダだけを追加で取り込んでも、ほかの root の登録は削除されない。"""
        other = os.path.join(self.tmp, "他のフォルダ")
        os.makedirs(other)
        import docx
        d = docx.Document()
        d.add_paragraph("追加フォルダの本文ゼータ")
        d.save(os.path.join(other, "追加.docx"))
        before = dump(self.conn)
        R.add_root(self.conn, other, self.cfg.roots)
        st = run_index(self.cfg, self.conn, roots=[other])
        R.clear_incomplete(self.conn, [other])
        self.assertEqual(st.deleted, [])
        self.assertEqual(dump(self.conn) - before, {("追加.docx", "段落1", "追加フォルダの本文ゼータ")})
        self.assertEqual(before - dump(self.conn), set())
        self.assertEqual(sorted(R.indexed_roots(self.conn)), sorted([self.docs, other]))
        self.assertEqual(self.names(other, "ゼータ"), ["追加.docx"])
        self.assertTrue(db.check_integrity(self.conn)[0])

    def test_cancelled_index_stays_incomplete_and_resumes(self):
        """インデックス作成を中止したフォルダは「未完了」のままで、選び直すと続きから再開できる。"""
        conn = db.connect_rw(os.path.join(self.tmp, "cancel.db"))
        self.addCleanup(conn.close)
        R.add_root(conn, self.docs, [])

        class Stop(Exception):
            pass

        def stop(done, total, elapsed, cur):
            if done == 3:
                raise Stop()
        with self.assertRaises(Stop):
            run_index(self.cfg, conn, stop, roots=[self.docs])
        self.assertEqual(R.indexed_roots(conn), [])
        self.assertEqual(R.scopes_for(self.docs, R.indexed_roots(conn)), ([], False))  # 未完了なので「登録済み」とは扱わない
        R.add_root(conn, self.docs, [])
        st = run_index(self.cfg, conn, roots=[self.docs])
        self.assertEqual(st.skipped_unchanged, 3)
        R.clear_incomplete(conn, [self.docs])
        self.assertEqual(dump(conn), dump(self.conn))

    def test_cli_add_and_folder_search(self):
        """index.py --add で、roots が空の設定でもフォルダを追加でき、search.py --folder で、そのフォルダ以下だけを検索できる。"""
        cfgfile = os.path.join(self.tmp, "c.toml")
        with open(cfgfile, "w", encoding="utf-8") as f:
            f.write('db_path = "%s"\noutput_dir = "%s"\n' % (os.path.join(self.tmp, "cli.db").replace("\\", "/"), os.path.join(self.tmp, "o").replace("\\", "/")))
        out, err = io.StringIO(), io.StringIO()
        self.assertEqual(index_cli.main(["--config", cfgfile], out, err), 2)  # 対象が1つもない
        self.assertIn("検索対象のフォルダがありません", err.getvalue())
        self.assertEqual(index_cli.main(["--config", cfgfile, "--add", os.path.join(self.tmp, "無い")], io.StringIO(), err), 4)
        self.assertIn("アクセスできません", err.getvalue())
        out, err = io.StringIO(), io.StringIO()
        self.assertEqual(index_cli.main(["--config", cfgfile, "--add", os.path.join(self.docs, "契約書")], out, err), 0, err.getvalue())
        out, err = io.StringIO(), io.StringIO()
        self.assertEqual(search_cli.main(["--config", cfgfile, "ABC", "--folder", os.path.join(self.docs, "契約書")], out, err), 0, err.getvalue())
        self.assertIn("該当 2 ファイル", out.getvalue())
        out, err = io.StringIO(), io.StringIO()
        self.assertEqual(search_cli.main(["--config", cfgfile, "ABC", "--folder", os.path.join(self.docs, "報告")], out, err), 2)
        self.assertIn("まだインデックスされていません", err.getvalue())
        out = io.StringIO()  # 追加後の通常の更新（引数なし）は、追加したフォルダを対象に動く
        self.assertEqual(index_cli.main(["--config", cfgfile], out, io.StringIO()), 0)
        self.assertIn("今回処理 0 ファイル", out.getvalue())


class ReadOnlyTest(unittest.TestCase):
    def test_target_folder_unchanged_after_everything(self):
        """インデックス作成・差分実行・検索・CSV保存・開く操作の前後で、対象フォルダの一覧・サイズ・更新日時が変わらない。"""
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        work, exp, cfg, conn = common.fresh_sample(tmp)
        self.addCleanup(conn.close)
        docs = cfg.roots[0]
        before = common.snapshot(docs)
        self.assertGreater(len(before), 30)
        writes = []
        real_open = builtins.open

        def spy(file, mode="r", *a, **k):
            if isinstance(file, (str, bytes, os.PathLike)) and os.fspath(file).startswith(docs) and set(str(mode)) & set("wax+"):
                writes.append((file, mode))
            return real_open(file, mode, *a, **k)
        old_tmp = tempfile.tempdir
        tempfile.tempdir = os.path.join(tmp, "ostmp")
        os.makedirs(tempfile.tempdir)
        try:
            with mock.patch.object(builtins, "open", spy), mock.patch.object(opener, "_launch"):
                run_index(cfg, conn)
                run_index(cfg, conn)
                run_index(cfg, conn, retry_errors=True)
                ro = db.connect_ro(cfg.db_path)
                for scope in ("file", "place"):
                    for q in ("ABC", "納期", "委託料 損害賠償", "損害賠償 -上限"):
                        for r in search(ro, q, build_options(limit=100, scope=scope)).results:
                            opener.open_file(r.fullpath, "copy")
                            opener.open_file(r.fullpath, "original")
                ro.close()
                cfgfile = os.path.join(tmp, "c.toml")
                with real_open(cfgfile, "w", encoding="utf-8") as f:
                    f.write('roots = ["%s"]\ndb_path = "%s"\noutput_dir = "%s"\n' % tuple(
                        p.replace("\\", "/") for p in (docs, cfg.db_path, os.path.join(tmp, "out"))))
                rc = search_cli.main(["--config", cfgfile, "ABC", "--csv", os.path.join(tmp, "out", "r.csv")], io.StringIO(), io.StringIO())
                self.assertEqual(rc, 0)
            opener.cleanup_temp()
        finally:
            tempfile.tempdir = old_tmp
        self.assertEqual(writes, [], "対象フォルダ内のファイルを書き込みモードで開いた")
        self.assertEqual(common.snapshot(docs), before)


class ReadOnlyDbTest(unittest.TestCase):
    def test_search_connection_is_read_only(self):
        """検索用のDB接続は読み取り専用で、書き込もうとするとエラーになる。"""
        r = common.shared_verification()
        ro = db.connect_ro(r["cfg"].db_path)
        try:
            with self.assertRaises(Exception):
                ro.execute("DELETE FROM chunks")
        finally:
            ro.close()


if __name__ == "__main__":
    unittest.main()
