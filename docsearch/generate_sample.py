"""ダミー文書群（Word・Excel・PowerPoint・PDF）と expected.json を生成する。実在の文書は一切使わない。

使い方: python generate_sample.py [--out sample_out] [--bulk 200]
  <out>/docs/          … ダミー文書（日本語名・全角スペース入りの名前・長いパス・リンク等を含む）
  <out>/expected.json  … 仕込んだ内容と、検索語ごとの期待される「ファイル名・場所」

expected.json の期待値は、仕込んだ全本文（registry）に対し、このファイル内の単純な実装
（NFKC＋小文字＋部分一致）で求めたもの。検索ツール本体（core/）のコードは使っていない。
"""
import argparse
import io
import json
import os
import random
import re
import shutil
import unicodedata
import zlib
from datetime import datetime

# 期待される issues の種別名（検索ツール本体の core/kinds.py と一致させる。テストで突き合わせる）
K = {"legacy": "対象外(旧形式)", "link": "対象外(リンク)", "password": "パスワード付き", "aes": "読込不可（AES暗号化）",
     "corrupt": "破損", "notext": "テキスト抽出不可", "formula": "数式キャッシュなし"}
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
FW = "　"  # 全角スペース


def _n(s):
    """oracle 用の簡易正規化（NFKC＋小文字＋空白連続を1つに）。検索ツール本体とは別実装。"""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", s).lower()).strip()


# ------------------------------------------------------------------ PDF（標準ライブラリだけで作る）
def _hex16(s):
    """文字列を UTF-16BE の16進文字列にする（PDFのCIDフォント用）。"""
    return "<" + s.encode("utf-16-be").hex() + ">"


def _tounicode(chars):
    """使う文字だけの ToUnicode CMap（コード=Unicode）を作る。pypdf がテキストを復元するために必要。"""
    cps = sorted({ord(c) for c in chars if ord(c) < 0x10000})
    blocks = []
    for i in range(0, len(cps), 100):
        part = cps[i:i + 100]
        blocks.append("%d beginbfchar\n%s\nendbfchar" % (len(part), "\n".join("<%04x> <%04x>" % (c, c) for c in part)))
    return ("/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            "/CMapName /Adobe-Identity-UCS def /CMapType 2 def\n"
            "1 begincodespacerange <0000> <FFFF> endcodespacerange\n" + "\n".join(blocks) +
            "\nendcmap CMapName currentdict /CMap defineresource pop end end").encode("ascii")


def make_pdf(pages):
    """PDFのバイト列を作る。pages の各要素は 文字列のリスト（行）か None（画像のみのページ）。

    日本語は、埋め込みなしの CID フォント（HeiseiKakuGo-W5 / UniJIS-UCS2-H）＋ToUnicode で表す。
    """
    chars = "".join("".join(p) for p in pages if p)
    objs = {}
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    page_ids = []
    nxt = 10
    for p in pages:
        pid, cid = nxt, nxt + 1
        nxt += 2
        page_ids.append(pid)
        if p is None:
            content = b"q 400 0 0 400 100 300 cm /Im1 Do Q"
        else:
            lines = ["BT /F1 14 Tf 72 760 Td 20 TL"]
            for ln in p:
                lines.append("%s Tj T*" % _hex16(ln))
            lines.append("ET")
            content = "\n".join(lines).encode("ascii")
        objs[cid] = b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"
        objs[pid] = (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents %d 0 R "
                     b"/Resources << /Font << /F1 3 0 R >> /XObject << /Im1 6 0 R >> >> >>" % cid)
    objs[2] = ("<< /Type /Pages /Count %d /Kids [%s] >>" % (len(page_ids), " ".join("%d 0 R" % i for i in page_ids))).encode()
    objs[3] = (b"<< /Type /Font /Subtype /Type0 /BaseFont /HeiseiKakuGo-W5 /Encoding /UniJIS-UCS2-H "
               b"/DescendantFonts [4 0 R] /ToUnicode 5 0 R >>")
    objs[4] = (b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /HeiseiKakuGo-W5 /DW 1000 "
               b"/CIDSystemInfo << /Registry (Adobe) /Ordering (Japan1) /Supplement 2 >> >>")
    tu = _tounicode(chars)
    objs[5] = b"<< /Length %d >>\nstream\n" % len(tu) + tu + b"\nendstream"
    img = zlib.compress(bytes([(x * 32 + y * 8) % 256 for y in range(8) for x in range(8)]))
    objs[6] = (b"<< /Type /XObject /Subtype /Image /Width 8 /Height 8 /ColorSpace /DeviceGray /BitsPerComponent 8 "
               b"/Filter /FlateDecode /Length %d >>\nstream\n" % len(img) + img + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for i in sorted(objs):
        offsets[i] = out.tell()
        out.write(b"%d 0 obj\n" % i + objs[i] + b"\nendobj\n")
    size = max(objs) + 1
    xref = out.tell()
    out.write(b"xref\n0 %d\n" % size)
    out.write(b"0000000000 65535 f \n")
    for i in range(1, size):
        out.write(("%010d 00000 n \n" % offsets[i]).encode() if i in offsets else b"0000000000 65535 f \n")
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref))
    return out.getvalue()


def encrypt_pdf(raw, user_password, declare_aes=False):
    """pypdf で RC4-128 暗号化したPDFを作る。declare_aes=True なら、暗号化辞書だけ AESV2 と宣言する。

    本物のAES暗号化PDFを作るには cryptography が必要（使わない方針）。そのため AES版は
    「AESと宣言しているダミー」で、読み込み時に pypdf が cryptography を要求する経路の確認用。
    """
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, NumberObject
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(raw)))
    w.encrypt(user_password, algorithm="RC4-128")
    if declare_aes:
        e = w._encrypt_entry
        std = DictionaryObject({NameObject("/AuthEvent"): NameObject("/DocOpen"), NameObject("/CFM"): NameObject("/AESV2"),
                                NameObject("/Length"): NumberObject(16)})
        e[NameObject("/V")] = NumberObject(4)
        e[NameObject("/R")] = NumberObject(4)
        e[NameObject("/CF")] = DictionaryObject({NameObject("/StdCF"): std})
        e[NameObject("/StmF")] = NameObject("/StdCF")
        e[NameObject("/StrF")] = NameObject("/StdCF")
    b = io.BytesIO()
    w.write(b)
    return b.getvalue()


# ------------------------------------------------------------------ Excel のキャッシュ値の仕込み
def patch_cached_values(data, sheet_xml, cells):
    """openpyxl は数式の計算結果（キャッシュ値）を書かないため、xlsx内のシートXMLに値を書き足す。

    cells: {"B4": ("6200000", False), "C4": ("ＡＢＣ検収", True)}（True は文字列の結果）。
    Excelで保存したファイルと同じ形（<f> の後に <v>）にする。
    """
    zin = zipfile_read(data)
    out = io.BytesIO()
    import zipfile
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for name, content in zin:
            if name == sheet_xml:
                x = content.decode("utf-8")
                for coord, (val, is_str) in cells.items():
                    m = re.search(r'<c r="%s"[^>]*>(<f>.*?</f>)<v\s*/>(?:</v>)?</c>|<c r="%s"[^>]*>(<f>.*?</f>)<v></v></c>' % (coord, coord), x)
                    if not m:
                        raise RuntimeError("キャッシュ値を書き足す対象のセル %s が見つかりません" % coord)
                    f = m.group(1) or m.group(2)
                    new = '<c r="%s" t="str">%s<v>%s</v></c>' % (coord, f, val) if is_str else \
                          '<c r="%s">%s<v>%s</v></c>' % (coord, f, val)
                    x = x[:m.start()] + new + x[m.end():]
                content = x.encode("utf-8")
            zo.writestr(name, content)
    return out.getvalue()


def zipfile_read(data):
    """zip のバイト列から (名前, 内容) の一覧を返す。"""
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [(n, z.read(n)) for n in z.namelist()]


# ------------------------------------------------------------------ 本体
class Sample:
    """ダミー文書を作りながら、仕込んだ本文（registry）と、期待される問題（issues）を記録する。"""

    def __init__(self, out):
        self.out = os.path.abspath(out)
        self.docs = os.path.join(self.out, "docs")
        self.registry = []   # {"file","location","text"}
        self.mtimes = {}     # 相対パス -> "YYYY-MM-DD"
        self.expected_issues = []  # {"file","kind","location"}
        self.expected_files = []   # 取り込み対象（拡張子）のファイルの相対パス
        self.symlinks = []

    def _path(self, rel):
        """相対パスから実パスを作り、フォルダを用意する。"""
        p = os.path.join(self.docs, *rel.split("/"))
        os.makedirs(os.path.dirname(p), exist_ok=True)
        return p

    def _write(self, rel, data, mtime=None):
        """バイト列を書き、更新日を設定する。"""
        p = self._path(rel)
        with open(p, "wb") as f:
            f.write(data)
        if mtime:
            self.mtimes[rel] = mtime
            ts = datetime.strptime(mtime, "%Y-%m-%d").replace(hour=12).timestamp()
            os.utime(p, (ts, ts))

    def reg(self, rel, loc, text):
        """仕込んだ本文を記録する（空は記録しない）。"""
        if _n(text):
            self.registry.append({"file": rel, "location": loc, "text": text})

    def docx(self, rel, items, mtime):
        """Wordを作る。items は ("p", 文字列) または ("t", 行のリスト)。段落番号は空段落も数える。"""
        import docx
        d = docx.Document()
        pn = tn = 0
        for kind, val in items:
            if kind == "p":
                pn += 1
                d.add_paragraph(val)
                self.reg(rel, "段落%d" % pn, val)
            else:
                tn += 1
                t = d.add_table(rows=len(val), cols=len(val[0]))
                for r, row in enumerate(val):
                    for c, v in enumerate(row):
                        t.cell(r, c).text = v
                        self.reg(rel, "表%d %d行%d列" % (tn, r + 1, c + 1), v)
        bio = io.BytesIO()
        d.save(bio)
        self._write(rel, bio.getvalue(), mtime)
        self.expected_files.append(rel)
        return bio.getvalue()

    def pptx(self, rel, slides, mtime):
        """PowerPointを作る。slides は {"texts":[...], "table":[[...]], "notes": "..."} の列。"""
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        for n, s in enumerate(slides, 1):
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            y = 0.3
            for t in s.get("texts", []):
                tb = sl.shapes.add_textbox(Inches(0.5), Inches(y), Inches(8), Inches(0.8))
                tb.text_frame.text = t
                y += 0.9
            if s.get("texts"):
                self.reg(rel, "スライド%d" % n, "\n".join(s["texts"]))
            if s.get("table"):
                rows = s["table"]
                gf = sl.shapes.add_table(len(rows), len(rows[0]), Inches(0.5), Inches(y), Inches(8), Inches(1.5))
                for r, row in enumerate(rows):
                    for c, v in enumerate(row):
                        gf.table.cell(r, c).text = v
                self.reg(rel, "スライド%d(表)" % n, "\n".join(" | ".join(r) for r in rows))
            if s.get("notes"):
                sl.notes_slide.notes_text_frame.text = s["notes"]
                self.reg(rel, "スライド%d(ノート)" % n, s["notes"])
        bio = io.BytesIO()
        prs.save(bio)
        self._write(rel, bio.getvalue(), mtime)
        self.expected_files.append(rel)

    def xlsx(self, rel, sheets, mtime, hidden=(), formulas=None, cached=None):
        """Excelを作る。sheets は {シート名: {セル: 値}}。formulas は {シート名: {セル: 数式}}、
        cached は {シート名: {セル: (値, 文字列か)}}（数式のキャッシュ値の仕込み）。"""
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, cells in sheets.items():
            ws = wb.create_sheet(name)
            for coord, v in cells.items():
                ws[coord] = v
                if isinstance(v, str):
                    self.reg(rel, "%s!%s" % (name, coord), v)
            for coord, f in (formulas or {}).get(name, {}).items():
                ws[coord] = f
            if name in hidden:
                ws.sheet_state = "hidden"
        bio = io.BytesIO()
        wb.save(bio)
        data = bio.getvalue()
        for name, cells in (cached or {}).items():
            idx = list(sheets).index(name) + 1
            data = patch_cached_values(data, "xl/worksheets/sheet%d.xml" % idx, cells)
            for coord, (val, is_str) in cells.items():
                if is_str:
                    self.reg(rel, "%s!%s" % (name, coord), val)
        self._write(rel, data, mtime)
        self.expected_files.append(rel)

    def pdf(self, rel, pages, mtime, password=None, declare_aes=False, register=True):
        """PDFを作る。pages は 行のリスト／None（画像のみ）。password が空文字なら空パスワードのRC4暗号化。"""
        raw = make_pdf(pages)
        if password is not None:
            raw = encrypt_pdf(raw, password, declare_aes)
        self._write(rel, raw, mtime)
        self.expected_files.append(rel)
        if register:
            for n, p in enumerate(pages, 1):
                if p:
                    self.reg(rel, "p.%d" % n, " ".join(p))

    def issue(self, rel, kind, location=""):
        """期待される issues（対象外・エラー等）を記録する。"""
        self.expected_issues.append({"file": rel, "kind": kind, "location": location})


def build_main(s):
    """検証用の主なダミー文書を作る。"""
    s.docx("契約書/業務委託契約書" + FW + "雛形.docx", [
        ("p", "業務委託契約書"),
        ("p", "第1条 委託業務の内容は、別紙記載のとおりとする。"),
        ("p", "第２条 本契約の有効期間は１２か月とする。"),
        ("p", "甲は乙に対し、委託料として金500万円を支払う。"),
        ("p", ""),
        ("p", "ＡＢＣ株式会社（以下「甲」という。）"),
        ("p", "秘密情報の取扱いについては別途定める。"),
        ("t", [["項目", "内容"], ["納期", "2024年3月31日"], ["検収", "ＡＢＣ社が実施する"]]),
    ], "2023-06-15")
    s.docx("契約書/秘密保持契約書.docx", [
        ("p", "秘密保持契約書"),
        ("p", "第1条 秘密情報とは、開示された技術上・営業上の情報をいう。"),
        ("p", "abc株式会社(以下「甲」という。)"),
        ("p", "第12条 有効期間は12か月とする。"),
        ("p", "委託料の支払は別途定める。"),
        ("p", "損害賠償の上限は委託料の範囲とする。"),
        ("p", "㈱ダミー商事は、①の義務を負う。"),
    ], "2022-09-30")
    s.pptx("企画" + FW + "資料/新規事業企画書.pptx", [
        {"texts": ["新規事業企画書", "ダミー社 企画部"]},
        {"texts": ["市場規模と競合の分析", "ABC社との協業を検討する"]},
        {"texts": ["スケジュール"], "table": [["工程", "時期"], ["設計", "第1四半期"], ["納期", "第2四半期"]],
         "notes": "委託料の見積もりは別途確認すること"},
        {"texts": ["リスクと対策", "損害賠償の取扱いは法務に確認する"]},
    ], "2024-05-10")
    s.xlsx("企画" + FW + "資料/予算表 2024.xlsx", {
        "予算": {"A1": "項目", "B1": "金額", "A2": "委託料", "B2": 5000000, "A3": "広告宣伝費", "B3": 1200000, "A4": "合計"},
        "進行表": {"A1": "工程", "B1": "納期", "A2": "設計", "B2": "2024-04-30", "A3": "ＡＢＣ社向け納期", "B3": "未定"},
        "非表示": {"A1": "非表示シートの語句"},
    }, "2024-08-20", hidden=("非表示",),
        formulas={"予算": {"B4": "=SUM(B2:B3)", "C4": '="ＡＢＣ検収"', "D4": "=SUM(B2:B3)*2"}},
        cached={"予算": {"B4": ("6200000", False), "C4": ("ＡＢＣ検収", True)}})
    s.issue("企画" + FW + "資料/予算表 2024.xlsx", K["formula"], "予算!D4")
    s.pdf("報告/調査報告書.pdf", [["調査報告書", "ABC株式会社に関する調査結果を報告する。"], ["納期は2024年3月とする。"],
                              ["損害賠償の取扱いに関する調査"]], "2022-11-03")
    s.pdf("報告/混在PDF.pdf", [["混在ドキュメントの本文ページ"], None, ["最後のページの本文"]], "2024-02-01")
    s.issue("報告/混在PDF.pdf", K["notext"], "p.2")
    s.pdf("報告/スキャン資料.pdf", [None, None], "2023-03-03")
    s.issue("報告/スキャン資料.pdf", K["notext"], "全ページ")
    s.pdf("報告/暗号化PDF(空パスワードRC4).pdf", [["空パスワードで暗号化されたPDFの本文"]], "2023-10-10", password="")
    s.pdf("報告/暗号化PDF(パスワード付き).pdf", [["読めないはずの本文"]], "2023-10-11", password="dummy-pass", register=False)
    s.issue("報告/暗号化PDF(パスワード付き).pdf", K["password"])
    s.pdf("報告/暗号化PDF(AES宣言ダミー).pdf", [["AESダミーの本文"]], "2023-10-12", password="", declare_aes=True, register=False)
    s.issue("報告/暗号化PDF(AES宣言ダミー).pdf", K["aes"])
    s.docx("報告/共通メモ.docx", [("p", "共通メモ（報告用）"), ("p", "報告フォルダのメモ。")], "2024-01-20")
    s.docx("企画" + FW + "資料/共通メモ.docx", [("p", "共通メモ（企画用）"), ("p", "企画フォルダのメモ。")], "2024-06-20")
    s.docx("文書/空の文書.docx", [("p", "")], "2024-01-01")
    deep = "/".join(("長いフォルダ名" * 6)[:40] + str(i) for i in range(7))
    s.docx(deep + "/長いパスの文書.docx", [("p", "長いパスの文書の本文です。")], "2024-03-03")
    # ---- 壊れたファイル・パスワード付き(推定)・旧形式・Officeの一時ファイル
    with open(os.path.join(s.docs, *"契約書/秘密保持契約書.docx".split("/")), "rb") as f:
        good = f.read()
    s._write("壊れた/壊れたWord.docx", good[:len(good) // 2], "2024-01-01")
    s.issue("壊れた/壊れたWord.docx", K["corrupt"])
    s._write("壊れた/壊れたExcel.xlsx", b"PK\x03\x04" + os.urandom(64), "2024-01-01")
    s.issue("壊れた/壊れたExcel.xlsx", K["corrupt"])
    s._write("壊れた/壊れたPowerPoint.pptx", "これはPowerPointではありません".encode("utf-8"), "2024-01-01")
    s.issue("壊れた/壊れたPowerPoint.pptx", K["corrupt"])
    s._write("壊れた/壊れたPDF.pdf", b"%PDF-1.4\n" + os.urandom(200), "2024-01-01")
    s.issue("壊れた/壊れたPDF.pdf", K["corrupt"])
    for e in ("docx", "xlsx", "pptx"):  # 暗号化Officeは OLE 形式。本物ではなく、署名だけを持つ疑似ファイル
        s._write("パスワード付き/保護された文書(疑似).%s" % e, OLE_MAGIC + b"\x00" * 504, "2024-01-01")
        s.issue("パスワード付き/保護された文書(疑似).%s" % e, K["password"])
    for e in ("doc", "xls", "ppt"):
        s._write("旧形式/旧様式.%s" % e, OLE_MAGIC + b"\x00" * 504, "2020-01-01")
        s.issue("旧形式/旧様式.%s" % e, K["legacy"])
    s._write("契約書/~$業務委託契約書" + FW + "雛形.docx", b"\x00" * 162, "2024-01-01")
    s._write("文書/メモ.txt", "対象外の拡張子".encode("utf-8"), "2024-01-01")
    # ---- リンク（辿らないこと）
    try:
        os.makedirs(os.path.join(s.docs, "リンク"), exist_ok=True)
        os.symlink(os.path.join(s.docs, "契約書", "業務委託契約書" + FW + "雛形.docx"), os.path.join(s.docs, "リンク", "雛形へのリンク.docx"))
        os.symlink(os.path.join(s.docs, "契約書"), os.path.join(s.docs, "リンク", "契約書フォルダへのリンク"), target_is_directory=True)
        s.symlinks = ["リンク/雛形へのリンク.docx", "リンク/契約書フォルダへのリンク"]
        for r in s.symlinks:
            s.issue(r, K["link"])
    except (OSError, NotImplementedError):
        s.symlinks = []  # 権限が無い環境（Windows等）では作らない


# 検索語ごとの期待値の定義。positives/negatives は手で書いた正規化前の語（検索式の解析には依存しない）
SEARCHES = [
    {"id": "S01", "note": "半角ABC（全角ＡＢＣ・小文字abcにも一致）", "query": "ABC", "positives": ["ABC"]},
    {"id": "S02", "note": "全角ＡＢＣで検索しても同じ", "query": "ＡＢＣ", "positives": ["ＡＢＣ"]},
    {"id": "S03", "note": "小文字abcで検索しても同じ", "query": "abc", "positives": ["abc"]},
    {"id": "S04", "note": "半角数字で検索 → 全角数字の本文にも一致", "query": "12か月", "positives": ["12か月"]},
    {"id": "S05", "note": "全角数字で検索 → 半角数字の本文にも一致", "query": "１２か月", "positives": ["１２か月"]},
    {"id": "S06", "note": "2文字（LIKE検索になる）", "query": "納期", "positives": ["納期"], "full_scan": True},
    {"id": "S07", "note": "複数ファイルに出現（3文字以上）", "query": "損害賠償", "positives": ["損害賠償"]},
    {"id": "S08", "note": "AND・同じ場所内", "query": "委託料 損害賠償", "positives": ["委託料", "損害賠償"], "scope": "place"},
    {"id": "S09", "note": "AND・ファイル内（場所が離れていても該当）", "query": "委託料 損害賠償", "positives": ["委託料", "損害賠償"], "scope": "file",
     "best_location": {"契約書/秘密保持契約書.docx": "段落6"}},
    {"id": "S10", "note": "除外（短い除外語）", "query": "損害賠償 -上限", "positives": ["損害賠償"], "negatives": ["上限"], "scope": "place"},
    {"id": "S11", "note": "フレーズ（全角・半角が同一視される）", "query": "\"有効期間は12か月\"", "positives": ["有効期間は12か月"]},
    {"id": "S12", "note": "フレーズ：語順・間隔が違うものは一致しない", "query": "\"有効期間 12か月\"", "positives": ["有効期間 12か月"]},
    {"id": "S13", "note": "存在しない語", "query": "存在しない語句ゼロ", "positives": ["存在しない語句ゼロ"]},
    {"id": "S14", "note": "ノート", "query": "見積もり", "positives": ["見積もり"]},
    {"id": "S15", "note": "PowerPointの表", "query": "第1四半期", "positives": ["第1四半期"]},
    {"id": "S16", "note": "非表示シート", "query": "非表示シートの語句", "positives": ["非表示シートの語句"]},
    {"id": "S17", "note": "260文字超のパスのファイル", "query": "長いパスの文書", "positives": ["長いパスの文書"]},
    {"id": "S18", "note": "数式セルのキャッシュ値（文字列）", "query": "ABC検収", "positives": ["ABC検収"]},
    {"id": "S19", "note": "㈱→(株)、①→1 の正規化", "query": "(株)ダミー商事は、1の義務", "positives": ["(株)ダミー商事は、1の義務"]},
    {"id": "S20", "note": "同名ファイルが別フォルダに2つ", "query": "共通メモ", "positives": ["共通メモ"]},
    {"id": "S21", "note": "拡張子の絞り込み", "query": "ABC", "positives": ["ABC"], "ext": ["pptx", "pdf"]},
    {"id": "S22", "note": "更新日の絞り込み（2024-04-01以降）", "query": "損害賠償", "positives": ["損害賠償"], "since": "2024-04-01"},
    {"id": "S23", "note": "フォルダの絞り込み", "query": "ABC", "positives": ["ABC"], "folder": "契約書"},
    {"id": "S24", "note": "空パスワードのRC4暗号化PDFは読める", "query": "空パスワードで暗号化", "positives": ["空パスワードで暗号化"]},
    {"id": "S25", "note": "2文字＋3文字以上のAND（LIKEは絞り込み後）", "query": "納期 ABC", "positives": ["納期", "ABC"], "scope": "place"},
    {"id": "S26", "note": "全て2文字の語のAND・ファイル内", "query": "納期 検収", "positives": ["納期", "検収"], "scope": "file", "full_scan": True},
    {"id": "S27", "note": "除外語が3文字以上", "query": "ABC -株式会社", "positives": ["ABC"], "negatives": ["株式会社"], "scope": "place"},
]


def compute_expected(s):
    """registry から、各検索語の期待値（ファイル名・場所の一覧）を求める。"""
    reg = [dict(r, norm=_n(r["text"])) for r in s.registry]
    out = []
    for sp in SEARCHES:
        pos = [_n(t) for t in sp["positives"]]
        neg = [_n(t) for t in sp.get("negatives", [])]
        scope = sp.get("scope", "file")

        def file_ok(f):
            """絞り込み条件（拡張子・フォルダ・更新日）に合うファイルか。"""
            if sp.get("ext") and f.rsplit(".", 1)[-1] not in sp["ext"]:
                return False
            if sp.get("folder") and not f.startswith(sp["folder"] + "/"):
                return False
            if sp.get("since") and s.mtimes[f] < sp["since"]:
                return False
            return True
        files = {}
        for r in reg:
            if file_ok(r["file"]):
                files.setdefault(r["file"], []).append(r)
        e = dict(sp)
        e["scope"] = scope
        if scope == "place":
            e["expect"] = sorted([r["file"], r["location"]] for fl in files.values() for r in fl
                                 if all(t in r["norm"] for t in pos) and not any(t in r["norm"] for t in neg))
        else:
            hit_files = sorted(f for f, fl in files.items()
                               if all(any(t in r["norm"] for r in fl) for t in pos)
                               and not any(t in r["norm"] for r in fl for t in neg))
            e["expect_files"] = hit_files
            e["expect_locations"] = {f: [r["location"] for r in files[f] if any(t in r["norm"] for t in pos)] for f in hit_files}
        out.append(e)
    return out


def generate(out_dir, quiet=False):
    """ダミー文書と expected.json を作り、expected（辞書）を返す。既存の out_dir は作り直す。"""
    out_dir = os.path.abspath(out_dir)
    if os.path.exists(out_dir):
        shutil.rmtree(out_dir)
    os.makedirs(out_dir)
    s = Sample(out_dir)
    build_main(s)
    kinds = {}
    for i in s.expected_issues:
        kinds[i["kind"]] = kinds.get(i["kind"], 0) + 1
    exp = {
        "about": "ダミー文書の仕込み内容と、検索語ごとの期待結果。location の書式は README を参照。",
        "root": "docs",
        "files": sorted(set(s.expected_files)),
        "mtimes": s.mtimes,
        "registry": s.registry,
        "expected_issues": s.expected_issues,
        "expected_issue_counts": kinds,
        "symlinks": s.symlinks,
        "pattern_excluded": 1,
        "searches": compute_expected(s),
    }
    with open(os.path.join(out_dir, "expected.json"), "w", encoding="utf-8") as f:
        json.dump(exp, f, ensure_ascii=False, indent=1)
    if not quiet:
        print("ダミー文書を生成しました: %s（取り込み対象 %d ファイル、仕込んだ本文 %d か所）" % (s.docs, len(exp["files"]), len(s.registry)))
    return exp


# ------------------------------------------------------------------ 規模・速度の目安を測るための大量生成
_PARTS = ["委託業務の範囲について定める", "本契約の有効期間は一年間とする", "納品物の検収は発注者が行う", "秘密情報を第三者に開示してはならない",
          "損害賠償の請求は直接損害に限る", "再委託には事前の書面による承諾を要する", "個人情報の取扱いは別紙のとおりとする",
          "料金の支払は月末締め翌月末払いとする", "本件に関する紛争は協議により解決する", "業務の進捗は週次で報告する",
          "仕様の変更は双方の合意により行う", "成果物の著作権は発注者に帰属する", "監査の実施について定める", "予算の範囲内で実施する"]


def generate_bulk(out_dir, n, seed=1):
    """速度・サイズ比の目安用に、ランダムな日本語文のWord/Excel/PowerPointを n 個作る（ダミー）。"""
    import docx
    import openpyxl
    from pptx import Presentation
    from pptx.util import Inches
    rnd = random.Random(seed)
    os.makedirs(out_dir, exist_ok=True)

    def para():
        """ダミーの1段落。"""
        return "。".join(rnd.choice(_PARTS) for _ in range(rnd.randint(2, 6))) + "。"
    for i in range(n):
        sub = os.path.join(out_dir, "フォルダ%02d" % (i % 10))
        os.makedirs(sub, exist_ok=True)
        kind = i % 4
        if kind in (0, 1):
            d = docx.Document()
            for _ in range(rnd.randint(20, 120)):
                d.add_paragraph(para())
            d.save(os.path.join(sub, "文書%04d.docx" % i))
        elif kind == 2:
            wb = openpyxl.Workbook()
            ws = wb.active
            for r in range(1, rnd.randint(50, 400)):
                ws.cell(r, 1, rnd.choice(_PARTS))
                ws.cell(r, 2, rnd.randint(1, 100000))
            wb.save(os.path.join(sub, "表%04d.xlsx" % i))
        else:
            p = Presentation()
            for _ in range(rnd.randint(5, 30)):
                sl = p.slides.add_slide(p.slide_layouts[6])
                sl.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame.text = para()
            p.save(os.path.join(sub, "資料%04d.pptx" % i))


def main():
    """コマンドライン入口。"""
    ap = argparse.ArgumentParser(description="ダミー文書の生成")
    ap.add_argument("--out", default="sample_out")
    ap.add_argument("--bulk", type=int, default=0, help="速度・サイズ比の目安用に、追加でN個の文書を <out>/bulk に作る")
    a = ap.parse_args()
    generate(a.out)
    if a.bulk:
        generate_bulk(os.path.join(a.out, "bulk"), a.bulk)
        print("大量生成: %d 個 → %s" % (a.bulk, os.path.join(a.out, "bulk")))


if __name__ == "__main__":
    main()
