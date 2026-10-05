"""Word・Excel・PowerPoint・PDF から、本文と「場所」を取り出す。

すべて、読み込み済みのバイト列（BytesIO）に対して処理する。元ファイルには触れない。
OCR は行わない。取り出せなかったものは ExtractError か issues として呼び出し側に返す。
"""
import io
import logging
import warnings
import zipfile
from dataclasses import dataclass, field

from .kinds import (KIND_AES, KIND_CELLCAP, KIND_CORRUPT, KIND_ERROR, KIND_FORMULA, KIND_NOTEXT, KIND_PASSWORD)
from .normalize import normalize

logging.getLogger("pypdf").setLevel(logging.ERROR)

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
MAX_ISSUES_PER_KIND = 50  # 1ファイル・1種別あたりに記録する issues の上限（超過分は件数のみ）


class ExtractError(Exception):
    """ファイル全体を読めなかった。kind は種別（kinds.py）、detail は利用者向けの原因と対処。"""

    def __init__(self, kind, detail):
        super().__init__(detail)
        self.kind = kind
        self.detail = detail


@dataclass
class Extracted:
    """抽出結果。chunks は [(場所, 正規化済み本文)]、issues は [(種別, 場所, 詳細)]。"""
    chunks: list = field(default_factory=list)
    issues: list = field(default_factory=list)
    truncated: bool = False

    def add(self, location, text):
        """本文を正規化して追加する（空なら追加しない）。"""
        t = normalize(text)
        if t:
            self.chunks.append((location, t))


class _IssueCollector:
    """種別ごとに issues の件数を制限して集める（大量発生によるDB肥大を防ぐ）。"""

    def __init__(self, out):
        self.out = out
        self.counts = {}

    def add(self, kind, location, detail):
        """issue を1件追加する。上限を超えたら数えるだけにする。"""
        n = self.counts.get(kind, 0) + 1
        self.counts[kind] = n
        if n <= MAX_ISSUES_PER_KIND:
            self.out.issues.append((kind, location, detail))

    def finish(self):
        """上限を超えた分を「ほか N 件」として1行にまとめる。"""
        for kind, n in self.counts.items():
            if n > MAX_ISSUES_PER_KIND:
                self.out.issues.append((kind, "(ほか%d件)" % (n - MAX_ISSUES_PER_KIND), "上限を超えた分は件数のみ記録しています"))


def _check_office_zip(data, label):
    """Office（OOXML）として読めるかを先に判定する。暗号化(OLE形式)・破損は ExtractError。"""
    if data[:8] == OLE_MAGIC:
        raise ExtractError(KIND_PASSWORD, "%s がパスワードで保護されている（暗号化されたOLE形式）可能性があります。"
                           "パスワードを解除したコピーを作成するか、対象外として扱ってください。" % label)
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExtractError(KIND_CORRUPT, "%s として読めません（ZIP構造が壊れている、または別形式のファイルです）。"
                           "元のアプリで開けるか確認してください。" % label)


def _generic_error(e, label):
    """想定外の例外を、種別つきの ExtractError に変換する。"""
    if isinstance(e, (zipfile.BadZipFile, KeyError)):
        return ExtractError(KIND_CORRUPT, "%s の内部構造が壊れています（%s: %s）。元のアプリで開けるか確認してください。"
                            % (label, type(e).__name__, e))
    return ExtractError(KIND_ERROR, "%s の読み込み中にエラーが発生しました（%s: %s）。"
                        "別のアプリで開けるか確認し、開けるのに失敗する場合は作成者に連絡してください。" % (label, type(e).__name__, e))


# ---------------------------------------------------------------- Word
def _docx_table(out, table, prefix):
    """Word の表のセルを「表N 行r列c」の場所で追加する。入れ子の表は再帰し、結合セルの重複は除く。"""
    seen = set()
    for r, row in enumerate(table.rows, 1):
        for c, cell in enumerate(row.cells, 1):
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            loc = "%s %d行%d列" % (prefix, r, c)
            out.add(loc, cell.text)
            for k, sub in enumerate(cell.tables, 1):
                _docx_table(out, sub, "%s/表%d" % (loc, k))


def extract_docx(data, cfg):
    """Word: 本文の段落（段落番号）と、表のセル（表番号・行・列）を取り出す。

    段落番号は本文直下の段落を、空の段落も含めて1から数える。ヘッダー・フッター・テキストボックス・脚注は対象外。
    """
    _check_office_zip(data, "Word文書")
    import docx
    out = Extracted()
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            doc = docx.Document(io.BytesIO(data))
            for i, p in enumerate(doc.paragraphs, 1):
                out.add("段落%d" % i, p.text)
            for t, table in enumerate(doc.tables, 1):
                _docx_table(out, table, "表%d" % t)
    except ExtractError:
        raise
    except Exception as e:
        raise _generic_error(e, "Word文書")
    return out


# ---------------------------------------------------------------- Excel
def _cell_text(cell, include_numbers):
    """セルの値から取り込む文字列を返す。取り込まないもの（数値・日付・エラー値）は None。"""
    v = cell.value
    if v is None or isinstance(v, bool):
        return None
    if getattr(cell, "data_type", "") == "e":
        return None
    if isinstance(v, str):
        return v
    if include_numbers and isinstance(v, (int, float)):
        return format(v, ".15g")
    return None


def extract_xlsx(data, cfg):
    """Excel: セルの文字列を「シート名!セル番地」の場所で取り出す。

    数式セルは計算結果（キャッシュ値）の文字列を対象にし、キャッシュが無い数式セルは issues に記録する。
    openpyxl の読み取り専用モードで、値用（data_only=True）と数式判定用（False）の2回、同時に走査する。
    文字列セルが cfg.max_cells_per_xlsx を超えたら、そこで打ち切る（走査セル数にも上限: 上限の10倍）。
    """
    _check_office_zip(data, "Excel文書")
    import openpyxl
    out = Extracted()
    issues = _IssueCollector(out)
    wb_v = wb_f = None
    cap = cfg.max_cells_per_xlsx
    scan_cap = cap * 10
    emitted = scanned = 0
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb_v = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            wb_f = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            stop = False
            for ws_v, ws_f in zip(wb_v.worksheets, wb_f.worksheets):
                for row_v, row_f in zip(ws_v.iter_rows(), ws_f.iter_rows()):
                    for cv, cf in zip(row_v, row_f):
                        scanned += 1
                        if scanned > scan_cap:
                            stop = True
                            break
                        if getattr(cf, "data_type", "") == "f" and cv.value is None:
                            issues.add(KIND_FORMULA, "%s!%s" % (ws_v.title, cf.coordinate),
                                       "数式のキャッシュ値がありません（Excelで開いて保存し直すと計算結果が保存されます）")
                            continue
                        text = _cell_text(cv, cfg.include_numbers_in_xlsx)
                        if text is None:
                            continue
                        if emitted >= cap:
                            stop = True
                            break
                        emitted += 1
                        out.add("%s!%s" % (ws_v.title, cv.coordinate), text)
                    if stop:
                        break
                if stop:
                    break
            if stop:
                out.truncated = True
                issues.add(KIND_CELLCAP, "", "セル数の上限（文字列セル %d 件／走査 %d セル）に達したため、以降は取り込んでいません。"
                           "config.toml の max_cells_per_xlsx で変更できます。" % (cap, scan_cap))
    except ExtractError:
        raise
    except Exception as e:
        raise _generic_error(e, "Excel文書")
    finally:
        for wb in (wb_v, wb_f):
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass
    issues.finish()
    return out


# ---------------------------------------------------------------- PowerPoint
def _pptx_walk(shapes, texts, tables):
    """スライド内の図形を（グループも再帰して）たどり、文字を texts、表を tables に集める。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            _pptx_walk(sh.shapes, texts, tables)
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = []
            for row in sh.table.rows:
                cells = [c.text for c in row.cells if not c.is_spanned]
                rows.append(" | ".join(cells))
            tables.append("\n".join(rows))
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            texts.append(sh.text_frame.text)


def extract_pptx(data, cfg):
    """PowerPoint: スライド内のテキスト（スライドN）・表（スライドN(表)）・ノート（スライドN(ノート)）を取り出す。

    スライド番号は並び順の1始まり。マスター・レイアウト、グラフ、SmartArt の文字は対象外。
    """
    _check_office_zip(data, "PowerPoint文書")
    from pptx import Presentation
    out = Extracted()
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            prs = Presentation(io.BytesIO(data))
            for n, slide in enumerate(prs.slides, 1):
                texts, tables = [], []
                _pptx_walk(slide.shapes, texts, tables)
                out.add("スライド%d" % n, "\n".join(texts))
                out.add("スライド%d(表)" % n, "\n".join(tables))
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                    out.add("スライド%d(ノート)" % n, slide.notes_slide.notes_text_frame.text)
    except ExtractError:
        raise
    except Exception as e:
        raise _generic_error(e, "PowerPoint文書")
    return out


# ---------------------------------------------------------------- PDF
def extract_pdf(data, cfg):
    """PDF: ページごとのテキストを「p.N」の場所で取り出す。

    テキストが取れないページ・ファイルは「テキスト抽出不可」として issues に記録する（OCRはしない）。
    パスワード付きは KIND_PASSWORD、AES暗号化（cryptography が無いと読めない）は KIND_AES の ExtractError。
    """
    from pypdf import PdfReader
    from pypdf.errors import DependencyError, EmptyFileError, PdfReadError
    out = Extracted()
    issues = _IssueCollector(out)
    aes = ExtractError(KIND_AES, "AESで暗号化されたPDFです。このツールは追加ライブラリ（cryptography）を使わない方針のため読めません。"
                                 "PDFの暗号化を外したコピーを作成してください。")
    empty_pages = []
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except DependencyError:
                raise aes
            if not ok:
                raise ExtractError(KIND_PASSWORD, "パスワード付きのPDFです。パスワードを解除したコピーを作成してください。")
        pages = reader.pages
        total = len(pages)
        if total == 0:
            raise ExtractError(KIND_CORRUPT, "PDFにページがありません（壊れている可能性があります）。")
        for n in range(total):
            try:
                text = pages[n].extract_text() or ""
            except DependencyError:
                raise aes
            except Exception as e:
                issues.add(KIND_ERROR, "p.%d" % (n + 1), "このページの読み込みに失敗しました（%s: %s）" % (type(e).__name__, e))
                continue
            t = normalize(text)
            if t:
                out.chunks.append(("p.%d" % (n + 1), t))
            else:
                empty_pages.append(n + 1)
        if empty_pages and len(empty_pages) == total:
            issues.add(KIND_NOTEXT, "全ページ", "%dページすべてでテキストを取得できません（画像のみのPDFの可能性。OCRは行いません）" % total)
        else:
            for n in empty_pages:
                issues.add(KIND_NOTEXT, "p.%d" % n, "このページはテキストを取得できません（画像のみの可能性。OCRは行いません）")
    except ExtractError:
        raise
    except DependencyError:
        raise aes
    except (PdfReadError, EmptyFileError) as e:
        raise ExtractError(KIND_CORRUPT, "PDFとして読めません（%s）。壊れている、または PDF ではないファイルです。" % e)
    except Exception as e:
        raise _generic_error(e, "PDF")
    issues.finish()
    return out


_EXTRACTORS = {"docx": extract_docx, "xlsx": extract_xlsx, "pptx": extract_pptx, "pdf": extract_pdf}


def extract(ext, data, cfg):
    """拡張子に応じた抽出関数を呼ぶ。失敗は ExtractError。"""
    return _EXTRACTORS[ext](data, cfg)
