"""テスト: python test_pdf2xlsx.py （テスト用に reportlab が必要。本体は不要）"""
import os
import subprocess
import sys
import tempfile
import unittest

import openpyxl
import pdfplumber

HERE = os.path.dirname(os.path.abspath(__file__))
S = os.path.join(HERE, "samples")
PDF = os.path.join(S, "sample_input.pdf")
TPL = os.path.join(S, "sample_template.xlsx")
EXP = os.path.join(S, "expected_output.xlsx")
TOOL = os.path.join(HERE, "pdf2xlsx.py")


def run_tool(pdf, out, tpl=TPL):
    env = dict(os.environ, PYTHONUTF8="1")  # 出力を文字列で検査するためUTF-8に固定
    p = subprocess.run([sys.executable, TOOL, pdf, tpl, out], capture_output=True,
                       text=True, encoding="utf-8", env=env)
    return p.returncode, p.stdout + p.stderr


def rebuild_pdf(src, dst, mutate):
    """src の単語を同じ座標で描き直した PDF を作る。mutate(page_no, words) で改変できる。"""
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    c = canvas.Canvas(dst, pagesize=landscape(A4))
    h = landscape(A4)[1]
    with pdfplumber.open(src) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            words = [dict(w) for w in page.extract_words()]
            words = mutate(n, words)
            for w in words:
                c.setFont("HeiseiKakuGo-W5", 9)
                c.drawString(w["x0"], h - w["bottom"] + 2, w["text"])
            c.showPage()
    c.save()


def drop_row(page, row_top):
    def f(n, words):
        if n != page:
            return words
        return [w for w in words if abs(w["top"] - row_top) > 3]
    return f


def second_detail_top(src, page=1):
    with pdfplumber.open(src) as pdf:
        tops = sorted({round(w["top"], 1) for w in pdf.pages[page - 1].extract_words()})
    return tops[4]  # タイトル2行+見出し行の次(1行目)の次 = 2明細目


class Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.t = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def test_matches_expected(self):
        out = os.path.join(self.t, "out.xlsx")
        rc, msg = run_tool(PDF, out)
        self.assertEqual(rc, 0, msg)
        got = openpyxl.load_workbook(out)["転記先"]
        exp = openpyxl.load_workbook(EXP)["転記先"]
        for r in range(1, 42):
            for col in "ABCDEFG":
                g, e = got[f"{col}{r}"], exp[f"{col}{r}"]
                self.assertEqual(g.value, e.value, f"{col}{r}")
                self.assertIs(type(g.value), type(e.value), f"{col}{r} 型")
                self.assertEqual(g.number_format, e.number_format, f"{col}{r} 書式")
        for cell in ("J2", "J3", "J4", "J5", "J6"):  # 数式は壊れていない／J4だけ値
            self.assertEqual(got[cell].value, exp[cell].value, cell)
        self.assertEqual(got["J4"].value, 6571658)

    def test_summary_values(self):
        """J2/J3/J6 の算出値（数式と同じ計算）が 40 / 6,571,658 / 一致 になる。"""
        out = os.path.join(self.t, "out.xlsx")
        run_tool(PDF, out)
        ws = openpyxl.load_workbook(out)["転記先"]
        col_a = [ws[f"A{r}"].value for r in range(2, 1001)]
        count = sum(v is not None for v in col_a)
        total = sum(ws[f"G{r}"].value or 0 for r in range(2, 1001))
        verdict = "一致" if total - ws["J4"].value == 0 else "不一致"
        self.assertEqual((count, total, verdict), (40, 6571658, "一致"))

    def test_input_not_modified(self):
        import hashlib
        def h(p):
            with open(p, "rb") as f:
                return hashlib.sha256(f.read()).hexdigest()
        before = (h(PDF), h(TPL))
        run_tool(PDF, os.path.join(self.t, "o.xlsx"))
        self.assertEqual(before, (h(PDF), h(TPL)))
        rc, _ = run_tool(PDF, TPL.replace(".xlsx", ".xlsx"), TPL)  # 出力=テンプレート
        self.assertNotEqual(rc, 0)

    def _expect_stop(self, mutate, keywords):
        bad = os.path.join(self.t, "bad.pdf")
        out = os.path.join(self.t, "bad.xlsx")
        rebuild_pdf(PDF, bad, mutate)
        rc, msg = run_tool(bad, out)
        self.assertNotEqual(rc, 0, msg)
        self.assertFalse(os.path.exists(out), "検証NGなのに出力ファイルができた")
        for k in keywords:
            self.assertIn(k, msg)

    def test_rebuilt_pdf_passes(self):
        """描き直しPDF（無改変）が通ること＝壊した版のNGが改変のせいだと言える対照実験。"""
        good = os.path.join(self.t, "good.pdf")
        rebuild_pdf(PDF, good, lambda n, w: w)
        rc, msg = run_tool(good, os.path.join(self.t, "g.xlsx"))
        self.assertEqual(rc, 0, msg)

    def test_row_deleted(self):
        self._expect_stop(drop_row(1, second_detail_top(PDF)), ["検証1", "検証2", "検証3", "1ページ"])

    def test_amount_changed(self):
        def f(n, words):
            if n == 2:
                for w in words:
                    if w["text"] == "202,000":
                        w["text"] = "202,001"
            return words
        self._expect_stop(f, ["検証1", "2ページ", "検証2"])

    def test_duplicate_voucher(self):
        def f(n, words):
            if n == 1:
                for w in words:
                    if w["text"] == "2026-0033":
                        w["text"] = "2026-0031"
            return words
        self._expect_stop(f, ["検証4"])

    def test_header_changed(self):
        def f(n, words):
            for w in words:
                if w["text"] == "支出先":
                    w["text"] = "相手先"
            return words
        self._expect_stop(f, ["検証5"])

    def test_bad_date(self):
        def f(n, words):
            if n == 1:
                for w in words:
                    if w["text"] == "令和8年10月01日":
                        w["text"] = "令和8年13月01日"
            return words
        self._expect_stop(f, ["検証6"])

    def test_no_detail_in_output(self):
        rc, msg = run_tool(PDF, os.path.join(self.t, "o.xlsx"))
        for s in ("テスト事務機", "2026-00", "資料印刷"):
            self.assertNotIn(s, msg)


if __name__ == "__main__":
    unittest.main(verbosity=2)
