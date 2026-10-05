"""検出結果と expected.json の一致、読み取り専用、再開、軽量モードなどの自動テスト
実行: python -m unittest discover -s tests -v   （folder_audit フォルダで）"""
from __future__ import annotations

import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import generate_sample  # noqa: E402
from audit import config as config_mod, hasher, pathutil, runner, scanner  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

WORK = None
EXP = None
CFG = None


def setUpModule():
    """ダミーツリーを生成し、設定を読み込む"""
    global WORK, EXP, CFG
    WORK = Path(tempfile.mkdtemp(prefix="audit_test_"))
    EXP = generate_sample.build(WORK)
    (WORK / "test_config.toml").write_text(generate_sample.TEST_CONFIG_TOML, encoding="utf-8")
    CFG = config_mod.load_config(WORK / "test_config.toml")


def tearDownModule():
    """権限を戻して一時フォルダを消す"""
    generate_sample.restore_permissions(WORK / "target")
    shutil.rmtree(WORK, ignore_errors=True)


def snapshot(target: Path):
    """対象フォルダ内の全エントリの (種別, サイズ, 更新日時ns, モード) を取る（読み取り専用の証明用）"""
    snap = {}
    for d, dirs, files in os.walk(target, followlinks=False):
        for n in dirs + files:
            p = os.path.join(d, n)
            st = os.lstat(p)
            snap[os.path.relpath(p, target)] = (st.st_mode, st.st_size, st.st_mtime_ns)
    return snap


def read_sheet(ws):
    """シートを (ヘッダー, [辞書]) で返す（3行目がヘッダー）"""
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    head = list(rows[0])
    return head, [dict(zip(head, r)) for r in rows[1:]]


class Base(unittest.TestCase):
    """通常モードを1回だけ実行して結果のExcelを共有する"""
    out = None
    wb = None
    before = None

    @classmethod
    def setUpClass(cls):
        """実行前のスナップショットを取り、ツールを実行する"""
        cls.target = WORK / "target"
        cls.out = Path(tempfile.mkdtemp(prefix="audit_out_", dir=WORK.parent))
        cls.before = snapshot(cls.target)
        res = runner.run_audit(str(cls.target), str(cls.out), CFG, out=io.StringIO())
        cls.res = res
        cls.after = snapshot(cls.target)
        cls.wb = load_workbook(res.xlsx_path)

    @classmethod
    def tearDownClass(cls):
        """出力フォルダを消す"""
        shutil.rmtree(cls.out, ignore_errors=True)

    def rel(self, folder, name):
        """フォルダ+ファイル名を対象からの相対POSIXパスにする"""
        return Path(os.path.relpath(os.path.join(folder, name), self.target)).as_posix()

    def sheet(self, name):
        """シートの行を辞書のリストで返す"""
        return read_sheet(self.wb[name])[1]


class TestDetection(Base):
    """検出結果が expected.json（仕込み）と一致するか"""

    def test_totals(self):
        """総ファイル数・総容量・除外数・リンク数"""
        s = {r["項目"].strip(): r for r in self.sheet("サマリ")}
        self.assertEqual(s["総ファイル数"]["値"], EXP["totals"]["files"])
        self.assertEqual(s["総ファイル数"]["容量(byte)"], EXP["totals"]["bytes"])
        self.assertEqual(s["除外ファイル数"]["値"], EXP["totals"]["excluded_files"])
        self.assertEqual(s["辿らなかったリンク等の件数"]["値"], EXP["totals"]["links"])
        self.assertEqual(s["フォルダ数"]["値"], EXP["totals"]["dirs"])
        self.assertEqual(s["0バイトのファイル数"]["値"], EXP["totals"]["zero_byte_files"])

    def test_duplicates(self):
        """全体一致グループ・未確認グループが一致し、誤検出がない"""
        groups = {}
        for r in self.sheet("重複候補"):
            groups.setdefault(r["グループID"], {"status": r["判定"], "size": r["1件あたりサイズ(byte)"],
                                              "wasted": r["グループの無駄容量(byte)"], "files": []})
            groups[r["グループID"]]["files"].append(self.rel(r["フォルダ"], r["ファイル名"]))
        conf = sorted(sorted(g["files"]) for g in groups.values() if g["status"] == hasher.CONFIRMED)
        unv = sorted(sorted(g["files"]) for g in groups.values() if g["status"] == hasher.UNVERIFIED)
        d = EXP["duplicates"]
        self.assertEqual(conf, sorted(g["files"] for g in d["confirmed_groups"]))
        self.assertEqual(unv, sorted(g["files"] for g in d["unverified_groups"]))
        self.assertEqual(sum(g["wasted"] for g in groups.values() if g["status"] == hasher.CONFIRMED),
                         d["confirmed_wasted_bytes"])
        detected = {f for g in groups.values() for f in g["files"]}
        for name, files in d["must_not_detect"].items():
            if name == "permission_pair_same_size" and EXP["unreadable"]["created"]:
                continue
            self.assertFalse(detected & set(files), f"誤検出: {name}")

    def test_naming(self):
        """必須・禁止パターンの逸脱が一致（過不足なし）"""
        rows = [r for r in self.sheet("命名逸脱") if r["区分"] == "命名規則"]
        req = {self.rel(r["フォルダ"], r["ファイル名"]) for r in rows if "必須パターン不一致" in r["該当ルール"]}
        forb = {self.rel(r["フォルダ"], r["ファイル名"]) for r in rows if "禁止パターン" in r["該当ルール"]}
        self.assertEqual(req, set(EXP["naming"]["required"]))
        self.assertEqual(forb, set(EXP["naming"]["forbidden"]))

    def test_version(self):
        """版管理パターン該当と、版の乱立グループが一致（別フォルダは束ねない）"""
        hits = {self.rel(r["フォルダ"], r["ファイル名"]) for r in self.sheet("命名逸脱") if r["区分"] == "版管理パターン"}
        self.assertEqual(hits, set(EXP["naming"]["version_hits"]))
        groups = {}
        for r in self.sheet("版の乱立候補"):
            groups.setdefault(r["グループID"], []).append(self.rel(r["フォルダ"], r["ファイル名"]))
        self.assertEqual(sorted(sorted(g) for g in groups.values()), sorted(EXP["version_groups"]))

    def test_stale(self):
        """長期未更新が一致し、境界当日は対象外"""
        got = {self.rel(r["フォルダ"], r["ファイル名"]) for r in self.sheet("長期未更新")}
        self.assertEqual(got, set(EXP["stale"]))
        self.assertFalse(got & set(EXP["not_stale_boundary"]))

    def test_folders(self):
        """フォルダ別の件数・容量・最古/最新が一致"""
        rows = self.sheet("フォルダ別")
        got = {}
        for r in rows:
            rel = Path(os.path.relpath(r["フォルダ"], self.target)).as_posix()
            got[rel] = r
        self.assertEqual(set(got), set(EXP["folders"]))
        for k, e in EXP["folders"].items():
            g = got[k]
            self.assertEqual((g["階層の深さ"], g["ファイル数"], g["合計容量(byte)"]), (e["depth"], e["files"], e["bytes"]), k)
            if e["min_mtime"] is None:
                self.assertIsNone(g["最も古い更新日"], k)
            else:
                import datetime as dt
                self.assertEqual(g["最も古い更新日"], dt.datetime.fromtimestamp(e["min_mtime"]).replace(microsecond=0), k)
                self.assertEqual(g["最も新しい更新日"], dt.datetime.fromtimestamp(e["max_mtime"]).replace(microsecond=0), k)

    def test_deep_folders(self):
        """階層が深いフォルダの一覧が一致"""
        got = {Path(os.path.relpath(r["フォルダ"], self.target)).as_posix() for r in self.sheet("階層が深いフォルダ")}
        self.assertEqual(got, set(EXP["deep_folders"]))

    def test_unreadable(self):
        """読み取り不可（作成できた環境のみ）。作れない環境では0件であること"""
        rows = self.sheet("読み取り不可")
        real = [r for r in rows if r["エラー種別"] != scanner.LINK_KIND]
        if EXP["unreadable"]["created"]:
            got = {Path(os.path.relpath(r["パス"], self.target)).as_posix() for r in real}
            self.assertEqual(got, set(EXP["unreadable"]["paths"]))
        else:
            self.assertEqual(real, [])
        links = [r for r in rows if r["エラー種別"] == scanner.LINK_KIND]
        self.assertEqual(len(links), EXP["totals"]["links"])

    def test_special_names(self):
        """長いパス・日本語・全角スペースのファイルが一覧に載る"""
        stale_or_dup = {self.rel(r["フォルダ"], r["ファイル名"]) for r in self.sheet("重複候補")}
        self.assertTrue(set(EXP["special_names"]["fullwidth_space_files"]) & stale_or_dup)
        self.assertGreater(EXP["special_names"]["long_path_relative_length"], 260)
        folders = {Path(os.path.relpath(r["フォルダ"], self.target)).as_posix() for r in self.sheet("フォルダ別")}
        self.assertIn(Path(EXP["special_names"]["long_path_file"]).parent.as_posix(), folders)


class TestSheetsFormat(Base):
    """シート構成・注意書き・フィルタ・ヘッダー固定"""

    def test_sheets(self):
        """必要なシートが揃い、注意書き・固定・フィルタがある"""
        names = ["サマリ", "重複候補", "命名逸脱", "版の乱立候補", "長期未更新", "フォルダ別", "階層が深いフォルダ", "読み取り不可"]
        self.assertEqual(self.wb.sheetnames, names)
        for ws in self.wb:
            self.assertEqual(ws["A1"].value, "内部資料のため取り扱い注意")
            self.assertEqual(ws.freeze_panes, "A4")
            self.assertTrue(ws.auto_filter.ref.startswith("A3:"))
        self.assertIn("更新日時は実際の作成・最終利用日を示さない場合があります", self.wb["長期未更新"]["A2"].value)

    def test_summary_settings(self):
        """サマリに設定値と基準日が出る"""
        items = {r["項目"].strip(): r["値"] for r in self.sheet("サマリ")}
        self.assertEqual(items["基準日"], "2026-10-01")
        self.assertEqual(items["stale_years"], 3)


class TestReadOnly(Base):
    """対象フォルダが変わっていないこと"""

    def test_unchanged(self):
        """実行前後でファイル一覧・サイズ・更新日時（とモード）が同一"""
        self.assertEqual(self.before, self.after)

    def test_no_work_files_in_target(self):
        """一時ファイル・出力が対象フォルダに作られない（出力は別フォルダのみ）"""
        self.assertTrue(any(p.suffix == ".xlsx" for p in self.out.iterdir()))
        self.assertFalse(list(self.target.rglob("*.sqlite")) + list(self.target.rglob("shared_folder_inventory_*")))

    def test_output_inside_target_refused(self):
        """出力先が対象の内側ならエラー"""
        with self.assertRaises(runner.RunError):
            runner.run_audit(str(self.target), str(self.target / "out"), CFG, out=io.StringIO())
        self.assertEqual(snapshot(self.target), self.before)


class TestModes(unittest.TestCase):
    """軽量モード・規則未設定・再開"""

    def setUp(self):
        """出力先を用意し、実行前の状態を取る"""
        self.target = WORK / "target"
        self.out = Path(tempfile.mkdtemp(prefix="audit_out2_", dir=WORK.parent))
        self.before = snapshot(self.target)

    def tearDown(self):
        """対象が不変であることを確認して出力を消す"""
        self.assertEqual(snapshot(self.target), self.before)
        shutil.rmtree(self.out, ignore_errors=True)

    def test_light_mode(self):
        """軽量モードはハッシュを計算せず、集計は通常モードと同じ"""
        with mock.patch.object(hasher, "partial_hash", side_effect=AssertionError("ハッシュ呼び出し")), \
                mock.patch.object(hasher, "full_hash", side_effect=AssertionError("ハッシュ呼び出し")):
            res = runner.run_audit(str(self.target), str(self.out), CFG, light=True, out=io.StringIO())
        wb = load_workbook(res.xlsx_path)
        self.assertEqual(read_sheet(wb["重複候補"])[1], [])
        self.assertIn("軽量モード", wb["重複候補"]["A2"].value)
        s = {r["項目"].strip(): r for r in read_sheet(wb["サマリ"])[1]}
        self.assertEqual(s["総ファイル数"]["値"], EXP["totals"]["files"])
        self.assertEqual(s["重複候補"]["値"], "軽量モードのため未実施")

    def test_no_naming_rules(self):
        """規則未設定でも動き、「規則未設定」と明記、必須/禁止の逸脱は0件"""
        cfg = config_mod.parse_config({"reference_date": "2026-10-01"})
        res = runner.run_audit(str(self.target), str(self.out), cfg, light=True, out=io.StringIO())
        wb = load_workbook(res.xlsx_path)
        self.assertIn("規則未設定", wb["命名逸脱"]["A2"].value)
        rows = read_sheet(wb["命名逸脱"])[1]
        self.assertFalse([r for r in rows if r["区分"] == "命名規則"])
        s = {r["項目"].strip(): r["値"] for r in read_sheet(wb["サマリ"])[1]}
        self.assertEqual(s["命名規則（必須/禁止パターン）"], "規則未設定")

    def test_resume(self):
        """全体ハッシュの途中で中断→--resume で再開し、結果が通常実行と一致、計算済みは再計算しない"""
        calls = {"n": 0}
        real = hasher.full_hash

        def interrupt(path):
            calls["n"] += 1
            if calls["n"] == 2:
                raise KeyboardInterrupt
            return real(path)

        with mock.patch.object(hasher, "full_hash", side_effect=interrupt):
            with self.assertRaises(KeyboardInterrupt):
                runner.run_audit(str(self.target), str(self.out), CFG, out=io.StringIO())
        self.assertTrue((self.out / runner.WORK_DB_NAME).exists())
        self.assertFalse(list(self.out.glob("*.xlsx")))
        counter = {"n": 0}

        def count(path):
            counter["n"] += 1
            return real(path)

        buf = io.StringIO()
        with mock.patch.object(hasher, "full_hash", side_effect=count):
            res = runner.run_audit(str(self.target), str(self.out), CFG, resume=True, out=buf)
        self.assertIn("再利用", buf.getvalue())
        # 全体ハッシュの対象は、A(3件)・試算A/B・同サイズA/B・権限ペア等。中断までに計算済みの分は再計算されない
        full_runs = {"n": 0}

        def count2(path):
            full_runs["n"] += 1
            return real(path)

        out2 = Path(tempfile.mkdtemp(prefix="audit_out3_", dir=WORK.parent))
        try:
            with mock.patch.object(hasher, "full_hash", side_effect=count2):
                ref = runner.run_audit(str(self.target), str(out2), CFG, out=io.StringIO())
        finally:
            pass
        self.assertLess(counter["n"], full_runs["n"])
        a = load_workbook(res.xlsx_path)
        b = load_workbook(ref.xlsx_path)
        for name in ("重複候補", "命名逸脱", "版の乱立候補", "長期未更新", "フォルダ別"):
            self.assertEqual(read_sheet(a[name]), read_sheet(b[name]), name)
        shutil.rmtree(out2, ignore_errors=True)

    def test_cli_end_to_end(self):
        """コマンドライン（main.py）で実行できる"""
        r = subprocess.run([sys.executable, str(ROOT / "main.py"), "--target", str(self.target), "--output",
                            str(self.out), "--config", str(WORK / "test_config.toml")],
                           capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("完了", r.stdout)
        self.assertTrue(list(self.out.glob("*.xlsx")))
        self.assertFalse((self.out / runner.WORK_DB_NAME).exists())


class TestErrorHandling(unittest.TestCase):
    """読み取れない項目で止まらず記録する（root環境でも検証できるようモックで再現）"""

    def test_scandir_and_hash_errors(self):
        """フォルダ列挙の権限エラー・パス長エラー・ハッシュ時の権限エラーを記録して続行"""
        target = WORK / "target"
        before = snapshot(target)
        out = Path(tempfile.mkdtemp(prefix="audit_out4_", dir=WORK.parent))
        real_scandir = os.scandir
        real_partial = hasher.partial_hash

        def fake_scandir(p):
            if p.endswith("総務部/規程"):
                raise PermissionError(13, "denied", p)
            if p.endswith("経理部/2020"):
                raise OSError(36, "File name too long", p)
            return real_scandir(p)

        def fake_partial(path, size, n):
            if path.endswith("提案資料_写し.pptx"):
                raise PermissionError(13, "denied", path)
            return real_partial(path, size, n)

        try:
            with mock.patch.object(scanner.os, "scandir", side_effect=fake_scandir), \
                    mock.patch.object(hasher, "partial_hash", side_effect=fake_partial):
                res = runner.run_audit(str(target), str(out), CFG, out=io.StringIO())
            wb = load_workbook(res.xlsx_path)
            rows = read_sheet(wb["読み取り不可"])[1]
            kinds = {}
            for r in rows:
                kinds.setdefault(r["エラー種別"], []).append(r["パス"])
            real_extra = len(EXP["unreadable"]["paths"])  # 非rootでは本物の権限エラーも加わる
            self.assertEqual(len(kinds["アクセス権限エラー"]), 2 + real_extra)   # モックのフォルダ1+ファイル1
            self.assertEqual(len(kinds["パス長エラー"]), 1)
            s = {r["項目"].strip(): r["値"] for r in read_sheet(wb["サマリ"])[1]}
            self.assertEqual(s["読み取り不可の件数"], 3 + real_extra)
            self.assertEqual(s["種類: アクセス権限エラー"], 2 + real_extra)
            # 失敗したファイルを除いた残り2件は重複候補として残る
            groups = [r for r in read_sheet(wb["重複候補"])[1] if r["フォルダ"].endswith("見積") or r["フォルダ"].endswith("配布用")]
            self.assertTrue(groups)
        finally:
            shutil.rmtree(out, ignore_errors=True)
        self.assertEqual(snapshot(target), before)


class TestUnits(unittest.TestCase):
    """パス変換・設定検証・日付計算の単体テスト"""

    def test_extended_path(self):
        """UNC・ドライブ・既に拡張形式のパスの変換（文字列のみ。Windows実機は未確認）"""
        f = lambda p: pathutil.to_extended_path(p, windows=True)
        self.assertEqual(f(r"\\srv\share\a\b"), "\\\\?\\UNC\\srv\\share\\a\\b")
        self.assertEqual(f("//srv/share/a"), "\\\\?\\UNC\\srv\\share\\a")
        self.assertEqual(f(r"C:\data\x"), "\\\\?\\C:\\data\\x")
        self.assertEqual(f("\\\\?\\C:\\data"), "\\\\?\\C:\\data")
        self.assertEqual(f("\\\\?\\UNC\\srv\\share"), "\\\\?\\UNC\\srv\\share")
        self.assertEqual(pathutil.to_extended_path("/a/b", windows=False), "/a/b")
        self.assertEqual(pathutil.to_display_path("\\\\?\\UNC\\srv\\share\\a"), "\\\\srv\\share\\a")
        self.assertEqual(pathutil.to_display_path("\\\\?\\C:\\a"), "C:\\a")

    def test_config_validation(self):
        """不正な正規表現・日付・上限値は分かりやすいエラーにする"""
        for bad in ({"naming": {"required_patterns": ["("]}}, {"reference_date": "2026/10/01"},
                    {"stale_years": 0}, {"hash": {"partial_bytes": 100, "max_read_bytes": 150}}):
            with self.assertRaises(config_mod.ConfigError):
                config_mod.parse_config(bad)

    def test_default_config_file_loads(self):
        """同梱の config.toml が読め、命名規則は未設定（規則なしで動く）"""
        cfg = config_mod.load_config(ROOT / "config.toml")
        self.assertEqual(cfg.stale_years, 3)
        self.assertEqual(cfg.depth_threshold, 8)
        self.assertEqual(cfg.required_patterns, ())

    def test_cutoff_leap_day(self):
        """2月29日の基準日でも境界日を計算できる"""
        import datetime as dt
        from audit.util import years_before
        self.assertEqual(years_before(dt.date(2028, 2, 29), 3), dt.date(2025, 2, 28))


if __name__ == "__main__":
    unittest.main()
