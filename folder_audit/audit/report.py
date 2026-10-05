"""Excel出力（1ファイル・複数シート）"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .util import safe_datetime

WARN = "内部資料のため取り扱い注意"
MTIME_NOTE = "更新日時は実際の作成・最終利用日を示さない場合があります（コピーや移動で変わるため）。"
MAX_DATA_ROWS = 1_048_576 - 3
DT_FMT = "yyyy-mm-dd hh:mm:ss"
HEADER_FILL = PatternFill("solid", fgColor="DDE6F0")


def _clean(v):
    """Excelに書けない文字を除き、長すぎる文字列を切り詰める"""
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:32767]
    return v


def write_sheet(wb, title, headers, rows, notes, formats=None, first=False):
    """1枚のシートを書く。1行目=注意書き、2行目=注記、3行目=ヘッダー（固定・フィルタ付き）"""
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    formats = formats or {}
    rows = list(rows)
    notes = list(notes)
    if len(rows) > MAX_DATA_ROWS:
        notes.append(f"※Excelの行数上限のため先頭{MAX_DATA_ROWS:,}行のみ表示（全{len(rows):,}行）")
        rows = rows[:MAX_DATA_ROWS]
    ws.cell(1, 1, WARN).font = Font(bold=True, color="C00000", size=12)
    ws.cell(2, 1, " / ".join(notes))
    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    widths = [len(str(h)) * 2 for h in headers]
    for r, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            v = _clean(v)
            cell = ws.cell(r, c, v)
            if isinstance(v, str):
                if v.startswith("="):
                    cell.data_type = "s"
                if "\n" in v:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if r < 200:
                    widths[c - 1] = max(widths[c - 1], min(len(v) * 2, 80))
            elif c in formats and v is not None:
                cell.number_format = formats[c]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = min(max(w, 8), 80)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, 3 + len(rows))}"
    return ws


def _d(ts):
    """UNIX時刻をExcel用の日時に変換する"""
    return safe_datetime(ts)


def write_report(path, ctx) -> None:
    """ctx（実行結果の辞書）からExcelを書き出す"""
    wb = Workbook()
    light = ctx["light"]
    # サマリ
    write_sheet(wb, "サマリ", ["区分", "項目", "値", "容量(byte)", "備考"], ctx["summary_rows"],
                [MTIME_NOTE], formats={3: "#,##0", 4: "#,##0"}, first=True)
    # 重複候補
    if light:
        rows, notes = [], ["軽量モードのためハッシュ計算を行わず、重複候補は未実施です"]
    else:
        rows = []
        for g in ctx["dups"]:
            for p in g["paths"]:
                st = ctx["file_info"](p)
                rows.append((g["id"], g["status"], len(g["paths"]), g["size"], g["wasted"],
                             os.path.dirname(p), os.path.basename(p), _d(st[1])))
        notes = ["「重複候補」です。配布用コピー等の正当な場合もあるため、重複と断定しません",
                 "0バイトのファイルは対象外です",
                 "「全体は未確認」は読み込み上限超過のため先頭・末尾のみ一致（全体の一致は未確認）",
                 "無駄容量=(件数−1)×サイズ。グループ内の各行に同じ値を表示（合計時は重複に注意）"]
    write_sheet(wb, "重複候補",
                ["グループID", "判定", "グループ内件数", "1件あたりサイズ(byte)", "グループの無駄容量(byte)",
                 "フォルダ", "ファイル名", "更新日時"], rows, notes,
                formats={4: "#,##0", 5: "#,##0", 8: DT_FMT})
    # 命名逸脱
    rows, notes = [], []
    if ctx["naming_state"] is None:
        notes.append("規則未設定（config の naming.required_patterns / forbidden_patterns が空のため、命名規則のチェックは行っていません）")
    for _p, d, n, why, size, m in ctx["naming"]:
        rows.append(("命名規則", why, d, n, size, _d(m)))
    for _p, d, n, pat, size, m in ctx["version_hits"]:
        rows.append(("版管理パターン", pat, d, n, size, _d(m)))
    write_sheet(wb, "命名逸脱", ["区分", "該当ルール", "フォルダ", "ファイル名", "サイズ(byte)", "更新日時"],
                rows, notes + ["区分「版管理パターン」は config の version.patterns に該当したファイルです"],
                formats={5: "#,##0", 6: DT_FMT})
    # 版の乱立候補
    rows = []
    for g in ctx["version_groups"]:
        for p, n, matched, size, m in g["items"]:
            rows.append((g["id"], g["dir"], g["key"], len(g["items"]), n, ", ".join(matched), size, _d(m)))
    write_sheet(wb, "版の乱立候補",
                ["グループID", "フォルダ", "正規化名", "グループ内件数", "ファイル名", "版パターン該当", "サイズ(byte)", "更新日時"],
                rows, ["同じフォルダ内で、版パターンを除いた名前と拡張子が同じファイルをまとめた「候補」です",
                       "版パターンに該当しない行は、同グループの元ファイルです"],
                formats={7: "#,##0", 8: DT_FMT})
    # 長期未更新
    rows = []
    ref = ctx["reference_date"]
    for d, n, size, m in ctx["stale"]:
        dtv = _d(m)
        years = round((ref - dtv.date()).days / 365.25, 1)
        rows.append((d, n, size, dtv, years))
    write_sheet(wb, "長期未更新", ["フォルダ", "ファイル名", "サイズ(byte)", "更新日時", "基準日からの経過年数"], rows,
                [f"基準日 {ref.isoformat()} の{ctx['cfg'].stale_years}年以上前（{ctx['cutoff_date'].isoformat()} 0時より前）に更新", MTIME_NOTE],
                formats={3: "#,##0", 4: DT_FMT})
    # フォルダ別
    rows = [(p, dep, n, s, _d(lo), _d(hi)) for p, dep, n, s, lo, hi in ctx["folders"]]
    write_sheet(wb, "フォルダ別", ["フォルダ", "階層の深さ", "ファイル数", "合計容量(byte)", "最も古い更新日", "最も新しい更新日"],
                rows, ["ファイル数・容量はそのフォルダ直下のファイルのみ（下位フォルダは含みません）。深さは対象フォルダ=0、直下のフォルダ=1", MTIME_NOTE],
                formats={4: "#,##0", 5: DT_FMT, 6: DT_FMT})
    # 階層が深いフォルダ
    rows = [(p, dep, n, s) for p, dep, n, s, _lo, _hi in ctx["deep"]]
    write_sheet(wb, "階層が深いフォルダ", ["フォルダ", "階層の深さ", "ファイル数", "合計容量(byte)"], rows,
                [f"階層の深さが {ctx['cfg'].depth_threshold} を超えるフォルダ"], formats={4: "#,##0"})
    # 読み取り不可
    write_sheet(wb, "読み取り不可", ["エラー種別", "パス", "詳細", "段階"], ctx["errors"],
                ["アクセス権限・パス長などで読み取れなかった項目、辿らなかったリンクです。黙って飛ばさず記録しています"])
    wb.save(path)
